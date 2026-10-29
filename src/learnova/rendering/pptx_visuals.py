"""
Native PowerPoint drawing for every visual family the Content Decision Engine can
choose. Shapes, tables and real chart objects — editable in PowerPoint, no
screenshots — so the .pptx matches the web deck instead of falling back to bullets.

Every drawer follows one grid (see ``layout``): content starts at the 0.5in margin,
cards are separated by the 0.22in gutter, and every card/box uses the same 0.22in
internal padding, icon disc and badge sizes. Type is fitted to the box (shared
across sibling cards so a row reads evenly) and compositions are centred in the band
so no visual leaves a big empty area or overflows.

Each drawer records ordered *groups* of shape ids (a card = rectangle + text + icon;
a timeline event; a table …). ``pptx_animation`` turns them into click builds.

``draw_family`` returns True if it drew something. Any family without a bespoke
drawing is rendered as a card grid built from its text, so a visual slide never
degrades to a plain bullet list.
"""

from __future__ import annotations

import io
import re
import math
from contextlib import contextmanager
from typing import Any, Callable, Dict, List, NamedTuple, Optional, Sequence, Tuple

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

from learnova.logging_config import logger
from learnova.rendering import layout as L
from learnova.visuals.icons import icon_png, pick_icons

PAD, GUT, ICON_D, BADGE_D = L.PAD, L.GUTTER, L.ICON_D, L.BADGE_D
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ── animation group registry ─────────────────────────────────────────────────
_GROUPS: Dict[int, List[List[int]]] = {}


def reset_groups() -> None:
    _GROUPS.clear()


def groups_for(slide) -> List[List[int]]:
    """Ordered shape-id groups recorded for ``slide`` (empty if none)."""
    return [list(g) for g in _GROUPS.get(slide.slide_id, [])]


def add_group(slide, ids: Sequence[int]) -> None:
    ids = [int(i) for i in ids]
    if ids:
        _GROUPS.setdefault(slide.slide_id, []).append(ids)


@contextmanager
def group(slide):
    """Record every shape created inside the block as one animation group."""
    n0 = len(slide.shapes)
    yield
    add_group(slide, [sh.shape_id for sh in list(slide.shapes)[n0:]])


# ── colour helpers ───────────────────────────────────────────────────────────
def _blend(rgb: RGBColor, t: float) -> RGBColor:
    """Mix ``rgb`` toward white by ``t`` (0 = rgb, 1 = white)."""
    r, g, b = rgb[0], rgb[1], rgb[2]
    return RGBColor(int(r + (255 - r) * t), int(g + (255 - g) * t), int(b + (255 - b) * t))


def _on(rgb: RGBColor) -> RGBColor:
    lum = 0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]
    return RGBColor(0x10, 0x14, 0x24) if lum > 150 else RGBColor(0xFF, 0xFF, 0xFF)


def _readable_accent(theme) -> RGBColor:
    """Accent colour for text on a white card — falls back to the primary when the accent is too pale."""
    a = theme.accent_rgb
    lum = 0.299 * a[0] + 0.587 * a[1] + 0.114 * a[2]
    return theme.primary_rgb if lum > 150 else a


def _hex(rgb: RGBColor) -> str:
    return "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])


class P(NamedTuple):
    """One paragraph: text, size (pt), bold, colour, space-before (pt), hanging bullet."""
    text: str
    size: float
    bold: bool = False
    color: Optional[RGBColor] = None
    before: float = 0.0
    bullet: bool = False
    level: int = 0                 # 0 = main point, 1 = sub-point (never deeper: research says two levels at most)


def _bullet(p, color: RGBColor, indent_in: float = 0.28, level: int = 0) -> None:
    """Hanging-indent bullet: wrapped lines align under the text, not under the glyph. Sub-points step in 0.32in with an en dash."""
    pPr = p._p.get_or_add_pPr()
    hang = indent_in if level == 0 else 0.24
    pPr.set("marL", str(int(Inches(indent_in + 0.32 * min(level, 1) + (hang - indent_in if level else 0)))))
    pPr.set("indent", str(-int(Inches(hang))))
    pPr.set("lvl", str(min(level, 1)))
    pPr.append(parse_xml(f'<a:buClr xmlns:a="{_A}"><a:srgbClr val="{_hex(color)}"/></a:buClr>'))
    pPr.append(parse_xml(f'<a:buFont xmlns:a="{_A}" typeface="Arial"/>'))
    pPr.append(parse_xml(f'<a:buChar xmlns:a="{_A}" char="{"•" if level == 0 else "–"}"/>'))


class _Canvas:
    def __init__(self, slide, theme, band, reserve: float = 0.0):
        self.s, self.t = slide, theme
        self.x, self.y = band.left, band.top
        self.w, self.h = band.width, band.height - reserve
        self.tint = _blend(theme.primary_rgb, 0.9)

    # -- grouping --------------------------------------------------------------
    def group(self):
        return group(self.s)

    # -- primitives --------------------------------------------------------------
    def shape(self, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.1):
        shp = self.s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(max(0.05, w)), Inches(max(0.05, h)))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                shp.adjustments[0] = min(0.5, radius)
            except Exception:
                pass
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(line_w)
        shp.shadow.inherit = False
        return shp

    def card(self, x, y, w, h, fill=None, line=None, radius=None):
        # radius is an absolute corner size so small and large cards share one visual radius;
        # unless a caller asks for its own, the deck template decides it (and whether cards have an outline)
        from learnova.rendering.templates import get_template

        tpl = get_template(getattr(self.t, "template_id", None))
        r_abs = tpl.card_radius if radius is None else radius
        r = r_abs / max(0.3, min(w, h)) * 2.2
        if fill is None:
            fill = self.t.card_bg_rgb if tpl.card_border else _blend(self.t.primary_rgb, 0.95)
        if line is None and tpl.card_border:
            line = _blend(self.t.primary_rgb, 0.72)
        return self.shape(MSO_SHAPE.ROUNDED_RECTANGLE if r_abs > 0 else MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, radius=r)

    def accent_bar(self, x, y, h, w=0.09, color=None):
        return self.shape(MSO_SHAPE.RECTANGLE, x, y, w, h, color or self.t.accent_rgb)

    def tb(self, x, y, w, h, paras: Sequence[P], anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, shape=None,
           inset: Tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0), bullet_color=None):
        """Text box (or text inside ``shape``) with explicit insets, so text edges sit exactly on the grid."""
        shp = shape or self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(v) for v in inset)
        room = (w if w > 0 else shp.width / 914400.0) - inset[0] - inset[2]
        shrink = L.word_scale([(q.text, q.size, q.bold) for q in paras], room)     # never cut a word across two lines
        first = True
        for para in paras:
            if not str(para.text).strip():
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            if para.before and not first:
                p.space_before = Pt(para.before)
            first = False
            p.alignment = align
            if para.bullet:
                _bullet(p, bullet_color or self.t.accent_rgb, level=para.level)
            r = p.add_run()
            r.text = str(para.text)
            r.font.size = Pt(round(para.size * shrink, 1))
            r.font.bold = para.bold
            r.font.color.rgb = para.color or self.t.text_rgb
        return shp

    def text(self, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
        """Back-compat: runs are ``(text, size, bold, colour)`` tuples."""
        paras = [P(t, s, b, col, 4 if i else 0) for i, (t, s, b, col) in enumerate(runs)]
        return self.tb(x, y, w, h, paras, anchor, align, shape, inset=(PAD * 0.5, 0.05, PAD * 0.5, 0.05) if shape else (0, 0, 0, 0))

    def fit(self, texts, w, h, mx=24.0, mn=10.0, spacing=6.0, bold=False) -> float:
        """Largest point size at which ``texts`` fill a w×h box without overflowing."""
        return L.fit_font_size([t for t in texts if t], max(0.3, w), max(0.2, h), max_pt=mx, min_pt=mn,
                               spacing_pt=spacing, bold=bold)

    def fit_paras(self, make: Callable[[float], List[P]], w, h, hi=24.0, lo=10.0) -> float:
        """Largest base size ``b`` (0.5pt steps) at which ``make(b)`` fits w×h; never below ``lo``."""
        b = hi
        while b > lo:
            ps = make(b)
            if L.paras_height([(p.text, p.size, p.bold, p.before) for p in ps], w) <= h and \
                    L.longest_word_fits([p.text for p in ps], w, min(p.size for p in ps), any(p.bold for p in ps)):
                return b
            b -= 0.5
        return lo

    def pair(self, head: str, body: str, b: float, hs: float = 1.2, head_color=None, body_color=None,
             lone_bold: Optional[bool] = None) -> List[P]:
        """Heading + body as paragraphs. A lone short phrase is a bold heading; a lone sentence is body text."""
        hc, bc = head_color or self.t.primary_rgb, body_color or self.t.text_rgb
        head, body = (head or "").strip(), (body or "").strip()
        if head and body:
            return [P(head, b * hs, True, hc), P(body, b, False, bc, 5)]
        t = head or body
        if lone_bold is None:
            lone_bold = len(t) <= 46
        if lone_bold:
            return [P(t, b * hs, True, hc)]
        return [P(t, b, False, bc)]

    def icon_circle(self, x, y, d, name, fill=None, color=None):
        """A tinted disc with a concept icon (PNG) centred in it."""
        fill = fill or _blend(self.t.primary_rgb, 0.86)
        self.shape(MSO_SHAPE.OVAL, x, y, d, d, fill)
        hexc = "#%02x%02x%02x" % tuple((color or self.t.primary_rgb)[i] for i in range(3))
        try:
            self.s.shapes.add_picture(io.BytesIO(icon_png(name, hexc, 192)), Inches(x + d * 0.2), Inches(y + d * 0.2),
                                      Inches(d * 0.6), Inches(d * 0.6))
        except Exception:
            pass

    def picture(self, png: bytes, x, y, w, h):
        """Place a PNG so it covers the (x, y, w, h) box without distortion (centre-cropped)."""
        from PIL import Image

        im = Image.open(io.BytesIO(png))
        iw, ih = im.size
        box_ratio, img_ratio = w / h, iw / ih
        if img_ratio > box_ratio:   # too wide → crop sides
            nw = int(ih * box_ratio)
            l = (iw - nw) // 2
            im = im.crop((l, 0, l + nw, ih))
        else:
            nh = int(iw / box_ratio)
            t = (ih - nh) // 2
            im = im.crop((0, t, iw, t + nh))
        buf = io.BytesIO()
        im.save(buf, "PNG")
        buf.seek(0)
        return self.s.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))

    def badge(self, x, y, n, d=BADGE_D, fill=None):
        fill = fill or self.t.primary_rgb
        shp = self.shape(MSO_SHAPE.OVAL, x, y, d, d, fill)
        self.tb(0, 0, 0, 0, [P(str(n), max(11, d * 26), True, _on(fill))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=shp)
        return shp

    # -- composition helpers ---------------------------------------------------
    def cells(self, n: int, cols: int, gut: float = GUT, max_ch: Optional[float] = None,
              area: Optional[Tuple[float, float, float, float]] = None):
        """Cards tiling the area: rows balanced, last row's cards widened to fill it, grid centred vertically."""
        ax, ay, aw, ah = area or (self.x, self.y, self.w, self.h)
        cols = max(1, min(cols, n))
        rows = math.ceil(n / cols)
        base, extra = divmod(n, rows)
        counts = [base + 1 if r < extra else base for r in range(rows)]
        ch = (ah - gut * (rows - 1)) / rows
        if max_ch and ch > max_ch:
            ch = max_ch
        oy = (ah - (ch * rows + gut * (rows - 1))) / 2
        out = []
        for r, cnt in enumerate(counts):
            cw = (aw - gut * (cnt - 1)) / cnt
            for k in range(cnt):
                out.append((ax + k * (cw + gut), ay + oy + r * (ch + gut), cw, ch))
        return out, ch


def _pt(n_chars: int, base: float, lo: float, scale: float = 1.0) -> float:
    return max(lo, min(base, base - (n_chars - 60) / 30.0)) * scale


def _flatten(data: Any, out: List[str]) -> None:
    if isinstance(data, str):
        if data.strip() and len(data) > 2:
            out.append(data.strip())
    elif isinstance(data, dict):
        for v in data.values():
            _flatten(v, out)
    elif isinstance(data, (list, tuple)):
        for v in data:
            _flatten(v, out)


def _same(a: str, b: str) -> bool:
    return a.strip().lower().rstrip(". ") == b.strip().lower().rstrip(". ")


def _lone_bold(items: Sequence[Tuple[str, str]]) -> bool:
    """Single-text cards are bold headings only when EVERY such card is a short phrase (one style per slide)."""
    singles = [(h or b) for h, b in items if not (h and b)]
    return bool(singles) and all(len(t) <= 46 for t in singles)


def _shared(fits: Sequence[float]) -> float:
    """Common type size for sibling cards: the smallest fit, but one outlier may not drag the rest below 70% of the median."""
    srt = sorted(fits)
    med = srt[len(srt) // 2]
    return max(min(fits), 0.7 * med)


def _auto_cols(n: int, w: float, min_cw: float = 3.3) -> int:
    fit = max(1, int((w + GUT) // (min_cw + GUT)))
    return max(1, min(n, fit, 4 if n == 4 and fit >= 4 else 3))


# ── card grid (icons / badges) ───────────────────────────────────────────────
def _grid(c: _Canvas, items: List[Tuple[str, str]], numbered: bool = False, cols: int = 0, icons: bool = True) -> None:
    """Cards tiling the canvas; each carries an icon (or badge) and text fitted to one shared size."""
    n = len(items)
    clean: List[Tuple[str, str]] = []
    for head, body in items:
        if body and (body.lower().startswith(head.lower().rstrip('. ')) or _same(head, body)):
            head, body = body, ""   # the heading was just the opening words — show the sentence once
        clean.append((head, body))
    lead = ICON_D + 0.16
    has_lead = icons or numbered
    lb = _lone_bold(clean)

    def plan(nc: int):
        cells, ch_full = c.cells(n, nc)
        horiz = ch_full <= 2.45 and min(w for _, _, w, _ in cells) >= 3.0

        def text_box(cell):
            x, y, w, h = cell
            if horiz:
                return w - 2 * PAD - ICON_D - 0.2, h - 0.16
            return w - 2 * PAD, h - 2 * PAD - (lead if has_lead else 0)
        fits = [c.fit_paras(lambda b, hd=hd, bd=bd: c.pair(hd, bd, b, lone_bold=lb), *text_box(cell), 26, 10.5)
                for (hd, bd), cell in zip(clean, cells)]
        return cells, ch_full, horiz, text_box, _shared(fits), fits

    if cols:
        cands = [cols]
    else:
        cands = [k for k in range(1, min(n, 4) + 1) if (c.w - GUT * (k - 1)) / k >= 2.7] or [1]
    best = None
    for k in cands:
        pl = plan(k)
        score = min(pl[5]) * 0.5 + pl[4] * 0.5 - (0.0 if n % math.ceil(n / k) == 0 or n % k == 0 else 0.8)
        if best is None or score > best[0] + 0.01:
            best = (score, k, pl)
    cols = best[1]
    cells, ch_full, horiz, text_box, b, fits = best[2]
    sz = [min(f, b) for f in fits]   # every card at the shared size; an outlier keeps its own smaller fit
    # shrink card height to the content when the content is much smaller than the band
    need = max(L.paras_height([(p.text, p.size, p.bold, p.before) for p in c.pair(h_, bd_, z, lone_bold=lb)], text_box(cells[0])[0])
               for (h_, bd_), z in zip(clean, sz))
    need_h = (need + 0.16 if horiz else need + 2 * PAD + (lead if has_lead else 0))
    max_ch = max(need_h * 1.28, ch_full * 0.62, 1.3) if need_h < ch_full else None
    if max_ch:
        cells, ch_full = c.cells(n, cols, max_ch=max_ch)
    used: List[str] = []
    for i, ((head, body), (x, y, cw, ch)) in enumerate(zip(clean, cells)):
        with c.group():
            c.card(x, y, cw, ch)
            if icons:
                mine = pick_icons(head + " " + body, 6)
                name = next((m for m in mine if m not in used), mine[0])
                used.append(name)
                if horiz:
                    c.icon_circle(x + PAD, y + (ch - ICON_D) / 2, ICON_D, name)
                else:
                    c.icon_circle(x + PAD, y + PAD, ICON_D, name)
            elif numbered:
                c.badge(x + PAD, y + (ch - BADGE_D) / 2 if horiz else y + PAD, i + 1)
            if horiz:
                tx = x + PAD + (ICON_D if icons else BADGE_D) + 0.2
                c.tb(tx, y + 0.08, x + cw - PAD - tx, ch - 0.16, c.pair(head, body, sz[i], lone_bold=lb), anchor=MSO_ANCHOR.MIDDLE)
            else:
                ty = y + PAD + (lead if has_lead else 0)
                c.tb(x + PAD, ty, cw - 2 * PAD, y + ch - PAD - ty, c.pair(head, body, sz[i], lone_bold=lb),
                     anchor=MSO_ANCHOR.TOP if has_lead else MSO_ANCHOR.MIDDLE)


# ── family drawers ────────────────────────────────────────────────────────────
def _process(c: _Canvas, data, cyclic: bool, foot_text: str = "") -> bool:
    raw = data.get("stages") or data.get("steps") or []
    items = []
    for st in raw[:8]:
        if isinstance(st, dict):
            items.append((str(st.get("label", "")), str(st.get("detail", ""))))
        else:
            items.append((str(st), ""))
    if len(items) < 2:
        return False
    if data.get("agenda") and len(items) >= 5:                   # a long agenda reads as a two-column numbered list, not tiny chevrons
        _grid(c, items, numbered=True, cols=2, icons=False)
        return True
    return _flow(c, items, cyclic, foot_text or ("↻  and the cycle repeats" if cyclic else ""), max_pt=20 if data.get("agenda") else 22)


def _flow(c: _Canvas, items: List[Tuple[str, str]], cyclic: bool = False, foot_text: str = "", max_pt: float = 22) -> bool:
    n = len(items)
    per = n if n <= 4 else math.ceil(n / 2)
    foot = 0.5 if foot_text else 0.0
    arrow_gap = 0.42
    area_h = c.h - foot
    rows = math.ceil(n / per)
    items = [(h, "" if _same(h, d) else d) for h, d in items]
    cw = (c.w - arrow_gap * (per - 1)) / per
    ch_full = (area_h - GUT * (rows - 1)) / rows

    horiz = ch_full < 2.4 or cw < 2.4

    def tbox(h_full):
        if horiz:
            return cw - 2 * PAD - BADGE_D - 0.18, h_full - 2 * PAD
        return cw - 2 * PAD, h_full - 2 * PAD - BADGE_D - 0.16
    lb = _lone_bold(items)
    fits = [c.fit_paras(lambda b, h_=h, d_=d: c.pair(h_, d_, b, lone_bold=lb), *tbox(ch_full), max_pt, 10.5) for h, d in items]
    b = _shared(fits)
    sz = [min(f, b) for f in fits]
    tw0 = tbox(ch_full)[0]
    need = max(L.paras_height([(p.text, p.size, p.bold, p.before) for p in c.pair(h, d, z, lone_bold=lb)], tw0) for (h, d), z in zip(items, sz))
    need_h = need + 2 * PAD + (0 if horiz else BADGE_D + 0.16)
    bare = not any(d for _, d in items)               # labels only (an agenda): cards hug the text instead of stretching
    ch = min(ch_full, max(need_h * 1.3, 1.4) if bare else max(need_h * 1.3, ch_full * 0.6, 1.4))
    total = rows * ch + GUT * (rows - 1) + foot
    oy = max(0.0, (c.h - total) / 2)
    for i, (lab, det) in enumerate(items):
        r, k = divmod(i, per)
        x, y = c.x + k * (cw + arrow_gap), c.y + oy + r * (ch + GUT)
        with c.group():
            if k > 0:   # connector arrow travels with the step it leads into
                c.shape(MSO_SHAPE.RIGHT_ARROW, x - arrow_gap + 0.06, y + ch / 2 - 0.17, arrow_gap - 0.12, 0.34, c.t.accent_rgb)
            c.card(x, y, cw, ch)
            if horiz:
                c.badge(x + PAD, y + (ch - BADGE_D) / 2, i + 1)
                c.tb(x + PAD + BADGE_D + 0.18, y + PAD, tw0, ch - 2 * PAD, c.pair(lab, det, sz[i], lone_bold=lb), anchor=MSO_ANCHOR.MIDDLE)
            else:
                c.badge(x + PAD, y + PAD, i + 1)
                c.tb(x + PAD, y + PAD + BADGE_D + 0.16, cw - 2 * PAD, ch - 2 * PAD - BADGE_D - 0.16, c.pair(lab, det, sz[i], lone_bold=lb))
    if foot_text:
        with c.group():
            c.tb(c.x, c.y + oy + total - foot + 0.06, c.w, foot - 0.06,
                 [P(foot_text, 16, True, _readable_accent(c.t))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
    return True


def _worked(c: _Canvas, data) -> bool:
    rows = data.get("rows") or [{"step": s, "reason": ""} for s in (data.get("steps") or [])]
    rows = [r for r in rows if str(r.get("step", "")).strip()][:8]
    if len(rows) < 2:
        return False
    n = len(rows)
    gap = GUT * 0.6
    rh_full = (c.h - gap * (n - 1)) / n
    tw = c.w - 3 * PAD - BADGE_D
    pairs = [(str(r["step"]), str(r.get("reason", "")).strip()) for r in rows]
    # step is the lead (bold), reason is the explanation
    def make(b, st, rs):
        ps = [P(st, b * 1.15, True, c.t.primary_rgb)]
        if rs:
            ps.append(P(rs, b, False, c.t.text_rgb, 3))
        return ps
    b = min(c.fit_paras(lambda s, st=st, rs=rs: make(s, st, rs), tw, rh_full - 0.16, 26, 10.5) for st, rs in pairs)
    need = max(L.paras_height([(p.text, p.size, p.bold, p.before) for p in make(b, st, rs)], tw) for st, rs in pairs)
    rh = min(rh_full, max(need + 0.3, 0.7, rh_full * 0.7))
    oy = (c.h - (rh * n + gap * (n - 1))) / 2
    for i, (st, rs) in enumerate(pairs):
        y = c.y + oy + i * (rh + gap)
        with c.group():
            c.card(c.x, y, c.w, rh)
            c.badge(c.x + PAD, y + (rh - BADGE_D) / 2, i + 1)
            c.tb(c.x + 2 * PAD + BADGE_D, y + 0.08, tw, rh - 0.16, make(b, st, rs), anchor=MSO_ANCHOR.MIDDLE)
    return True


_SUP = str.maketrans("0123456789+-n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ⁿ")


def _superscripts(text: str) -> str:
    """``x^2`` → ``x²`` (a caret is source-code notation, not something to show on a slide)."""
    return re.sub(r"\^\(?([0-9+\-n]{1,3})\)?", lambda m: m.group(1).translate(_SUP), text)


def _stepper(c: _Canvas, data) -> bool:
    """A solved problem: each step a numbered row (title + working, then the plain-language reason), the answer in a band below."""
    steps = [x for x in (data.get("steps") or []) if x.get("title")]
    if len(steps) < 2:
        return False
    if len(steps) > 6:
        steps = steps[:5] + steps[-1:]
    rows = []
    for st in steps:
        work = [m for m in (st.get("math") or []) if m][-1:] or [""]
        work_s = _superscripts(re.sub(r"\[\[|\]\]", "", work[0]))
        reason = str(st.get("explain", ""))
        if len(steps) >= 5:                                     # a long solution: the lead sentence only, so the type can stay readable
            reason = re.split(r"(?<=[.!?])\s+(?=[A-Z])", reason, maxsplit=1)[0]
            if len(reason) > 105:
                reason = reason[:105].rsplit(" ", 1)[0].rstrip(",;:") + "…"
        rows.append({"step": f"{st['title']}  —  {work_s}" if work_s and len(work_s) <= 48 else str(st["title"]), "reason": reason[:170]})
    band_h = 0.7
    h0 = c.h
    c.h = h0 - band_h - GUT * 0.6
    ok = _worked(c, {"rows": rows})
    c.h = h0
    if ok and data.get("answer"):
        with c.group():
            box = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, c.x, c.y + h0 - band_h, c.w, band_h, c.t.primary_rgb, radius=0.2)
            sz = c.fit([f"Answer: {data['answer']}"], c.w - 2 * PAD, band_h - 0.12, 22, 12, bold=True)
            c.tb(0, 0, 0, 0, [P(f"Answer: {data['answer']}", sz, True, _on(c.t.primary_rgb))], MSO_ANCHOR.MIDDLE, PP_ALIGN.LEFT, shape=box, inset=(PAD, 0.04, PAD, 0.04))
    return ok


def _timeline(c: _Canvas, data) -> bool:
    ev = [e for e in (data.get("events") or []) if e.get("title")][:7]
    if len(ev) < 3:
        return False
    n = len(ev)
    cy = c.y + c.h / 2
    step = c.w / n
    cw = min(step * 1.95, 3.7)
    ch = (c.h / 2) - 0.4
    ch = min(ch, 2.9)
    tw, th = cw - 2 * PAD, ch - 2 * PAD
    def make(b, e): return [P(str(e.get("date", "")), b * 1.35, True, _readable_accent(c.t)), P(str(e.get("title", "")), b, False, c.t.text_rgb, 3)]
    b = min(c.fit_paras(lambda s, e=e: make(s, e), tw, th, 24, 11) for e in ev)
    with c.group():
        c.shape(MSO_SHAPE.RECTANGLE, c.x, cy - 0.03, c.w, 0.06, _blend(c.t.primary_rgb, 0.4))
    for i, e in enumerate(ev):
        cx = c.x + step * (i + 0.5)
        up = i % 2 == 0
        x = min(max(c.x, cx - cw / 2), c.x + c.w - cw)
        y = cy - 0.4 - ch if up else cy + 0.4
        with c.group():
            c.shape(MSO_SHAPE.RECTANGLE, cx - 0.015, (y + ch) if up else cy, 0.03, 0.4 if up else (y - cy), _blend(c.t.primary_rgb, 0.4))
            c.shape(MSO_SHAPE.OVAL, cx - 0.14, cy - 0.14, 0.28, 0.28, c.t.accent_rgb, c.t.primary_rgb, 2.25)
            c.card(x, y, cw, ch)
            c.tb(x + PAD, y + PAD, tw, th, make(b, e), anchor=MSO_ANCHOR.MIDDLE)
    return True


def draw_table(c: _Canvas, heads: Sequence[Any], rows: Sequence[Sequence[Any]]) -> bool:
    heads, rows = list(heads), [list(r) for r in rows][:8]
    if len(heads) < 2 or len(rows) < 1:
        return False
    ncol = len(heads)
    nrow = len(rows) + 1
    cell_w = c.w / ncol
    pad_x = 0.16
    longest_per_row = [max((str(x) for x in r), key=len, default="") for r in rows]
    # the table fills the band but rows never exceed a comfortable height
    head_h = min(0.75, c.h / nrow)
    rh = min(1.15, (c.h - head_h) / len(rows))
    body_pt = min(c.fit([lg], cell_w - 2 * pad_x, rh - 0.14, 20, 10) for lg in longest_per_row)
    # tall table cells (wrapping) may need more height than rh — grow rows to what the text needs
    need = [max(L.estimate_lines(str(x), cell_w - 2 * pad_x, body_pt) for x in r) * body_pt * 1.3 / 72 + 0.22 for r in rows]
    heights = [max(rh, nd) for nd in need]
    total = head_h + sum(heights)
    if total > c.h:   # scale back down if growing overshot the band
        k = (c.h - head_h) / sum(heights)
        heights = [h * k for h in heights]
        total = c.h
    oy = (c.h - total) / 2
    with c.group():
        gf = c.s.shapes.add_table(nrow, ncol, Inches(c.x), Inches(c.y + oy), Inches(c.w), Inches(total))
        tbl = gf.table
        tbl.rows[0].height = Inches(head_h)
        for i, h in enumerate(heights):
            tbl.rows[i + 1].height = Inches(h)
        head_pt = min(20, max(12, body_pt + 3))
        for j, h in enumerate(heads):
            cell = tbl.cell(0, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = c.t.primary_rgb
            cell.margin_left = cell.margin_right = Inches(pad_x)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = str(h)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.size = Pt(head_pt)
                    r.font.bold = True
                    r.font.color.rgb = _on(c.t.primary_rgb)
        for i, row in enumerate(rows, start=1):
            for j in range(ncol):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = c.t.card_bg_rgb if i % 2 else _blend(c.t.primary_rgb, 0.92)
                cell.margin_left = cell.margin_right = Inches(pad_x)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text = str((row + [""] * ncol)[j])
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for r in p.runs:
                        r.font.size = Pt(body_pt)
                        r.font.bold = j == 0
                        r.font.color.rgb = c.t.primary_rgb if j == 0 else c.t.text_rgb
    return True


def _compare_table(c: _Canvas, data) -> bool:
    return draw_table(c, data.get("headers") or [], data.get("rows") or [])


def _pros_cons(c: _Canvas, data) -> bool:
    pros, cons = [str(x) for x in data.get("pros") or []][:6], [str(x) for x in data.get("cons") or []][:6]
    if not (pros and cons):
        return False
    cw = (c.w - GUT) / 2
    head_h = 0.55
    tw = cw - 2 * PAD - 0.3
    th = c.h - 2 * PAD - head_h - 0.15
    cols = (("Advantages", pros, RGBColor(0x1A, 0x7F, 0x4B)), ("Trade-offs", cons, RGBColor(0xB2, 0x3B, 0x3B)))
    def make(b, items, col): return [P(t, b, False, c.t.text_rgb, 8, True) for t in items]
    b = min(c.fit_paras(lambda s, it=it, col=col: make(s, it, col), tw, th, 24, 10.5) for _, it, col in cols)
    need = max(L.paras_height([(p.text, p.size, p.bold, p.before) for p in make(b, it, col)], tw) for _, it, col in cols)
    ch_all = min(c.h, max(need + 2 * PAD + head_h + 0.15 + 0.3, c.h * 0.6))
    oy = (c.h - ch_all) / 2
    th = ch_all - 2 * PAD - head_h - 0.15
    for k, (title, items, col) in enumerate(cols):
        x = c.x + k * (cw + GUT)
        with c.group():
            c.card(x, c.y + oy, cw, ch_all, line=col)
            hd = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PAD, c.y + oy + PAD, cw - 2 * PAD, head_h, col, radius=0.25)
            c.tb(0, 0, 0, 0, [P(title.upper(), 15, True, RGBColor(255, 255, 255))], MSO_ANCHOR.MIDDLE, PP_ALIGN.LEFT, shape=hd,
                 inset=(0.18, 0.02, 0.1, 0.02))
            c.tb(x + PAD, c.y + oy + PAD + head_h + 0.15, cw - 2 * PAD, th, make(b, items, col), anchor=MSO_ANCHOR.TOP, bullet_color=col)
    return True


def _references(c: _Canvas, data) -> bool:
    items = [str(x) for x in (data.get("items") or []) if str(x).strip()][:12]
    if not items:
        return False
    cols = 2 if len(items) > 4 else 1
    cells, ch_full = c.cells(len(items), cols, gut=GUT * 0.7)
    tw = min(w for _, _, w, _ in cells) - 3 * PAD - BADGE_D * 0.8
    d = min(BADGE_D, ch_full - 0.2)
    b = min(c.fit_paras(lambda s, t=t: [P(t, s, False, c.t.text_rgb)], tw, ch_full - 0.2, 24, 10) for t in items)
    need = max(L.paras_height([(t, b, False, 0)], tw) for t in items) + 0.3
    if need < ch_full * 0.7:
        cells, ch_full = c.cells(len(items), cols, gut=GUT * 0.7, max_ch=max(need * 1.3, 0.8))
    for i, (t, (x, y, cw, ch)) in enumerate(zip(items, cells)):
        with c.group():
            c.card(x, y, cw, ch)
            c.badge(x + PAD, y + (ch - d) / 2, i + 1, d=d)
            c.tb(x + 2 * PAD + d, y + 0.08, cw - 3 * PAD - d, ch - 0.16, [P(t, b, False, c.t.text_rgb)], anchor=MSO_ANCHOR.MIDDLE)
    return True


def _matrix(c: _Canvas, data) -> bool:
    q = (data.get("quadrants") or data.get("cells") or [])[:4]
    if len(q) < 3:
        return False
    cw, ch = (c.w - GUT) / 2, (c.h - GUT) / 2
    tw, th = cw - 2 * PAD, ch - 2 * PAD
    def make(b, qd):
        items = [str(t) for t in (qd.get("items") or [])][:4]
        return [P(str(qd.get("title", "")), b * 1.2, True, c.t.primary_rgb)] + [P(t, b, False, c.t.text_rgb, 4, True) for t in items]
    b = min(c.fit_paras(lambda s, qd=qd: make(s, qd), tw, th, 28, 10.5) for qd in q)
    for i, qd in enumerate(q):
        r, k = divmod(i, 2)
        x, y = c.x + k * (cw + GUT), c.y + r * (ch + GUT)
        with c.group():
            c.card(x, y, cw, ch, fill=_blend(c.t.primary_rgb, 0.9 - 0.04 * i))
            c.accent_bar(x, y + PAD, ch - 2 * PAD, 0.09)
            c.tb(x + PAD, y + PAD, tw, th, make(b, qd), anchor=MSO_ANCHOR.MIDDLE)
    return True


def _layers(c: _Canvas, data) -> bool:
    lv = (data.get("levels") or [])[:6]
    if len(lv) < 3:
        return False
    n = len(lv)
    rows = []
    for it in lv:
        label, detail = (str(it.get("label", "")), str(it.get("detail", ""))) if isinstance(it, dict) else (str(it), "")
        rows.append((label, "" if _same(label, detail) else detail))
    gap = GUT * 0.6
    rh_full = (c.h - gap * (n - 1)) / n
    narrow = 0.18
    min_w = c.w * (1 - narrow)
    has_detail = any(d for _, d in rows)
    lw = min(3.6, min_w * 0.32) if has_detail else min_w - 2 * PAD
    dw = min_w - lw - 3 * PAD
    sl = c.fit_paras(lambda s: [P(max(rows, key=lambda r: len(r[0]))[0], s, True)], lw, rh_full - 0.16, 24, 11)
    sd = c.fit_paras(lambda s: [P(max(rows, key=lambda r: len(r[1]))[1], s, False)], dw, rh_full - 0.16, 20, 10.5) if has_detail else 0
    rh = min(rh_full, 1.25)
    oy = (c.h - (rh * n + gap * (n - 1))) / 2
    for i, (label, detail) in enumerate(rows):
        frac = 1.0 - i * (narrow / max(1, n - 1))
        w = c.w * frac
        x = c.x + (c.w - w) / 2
        y = c.y + oy + i * (rh + gap)
        fill = _blend(c.t.primary_rgb, 0.8 - 0.8 * (i / max(1, n - 1)))
        fg = _on(fill)
        with c.group():
            c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, rh, fill, radius=0.18 / max(0.4, rh) * 0.3)
            c.tb(x + PAD, y + 0.06, lw if has_detail else w - 2 * PAD, rh - 0.12, [P(label, sl, True, fg)], anchor=MSO_ANCHOR.MIDDLE)
            if detail:
                c.tb(x + PAD + lw + PAD, y + 0.06, w - lw - 3 * PAD, rh - 0.12, [P(detail, sd, False, fg)], anchor=MSO_ANCHOR.MIDDLE)
    return True


def _venn(c: _Canvas, data) -> bool:
    items = [str(x) for x in data.get("items") or []][:8]
    if len(items) < 3:
        return False
    d = min(c.h * 0.92, c.w * 0.4)
    top = c.y + (c.h - d) / 2
    with c.group():
        c.shape(MSO_SHAPE.OVAL, c.x, top, d, d, _blend(c.t.primary_rgb, 0.7), c.t.primary_rgb, 1.5)
        c.shape(MSO_SHAPE.OVAL, c.x + d * 0.58, top, d, d, _blend(c.t.accent_rgb, 0.7), c.t.accent_rgb, 1.5)
    px = c.x + d * 1.58 + GUT
    pw = c.x + c.w - px
    tw, th = pw - 2 * PAD, c.h - 2 * PAD
    b = c.fit_paras(lambda s: [P(t, s, False, c.t.text_rgb, 8, True) for t in items], tw, th, 22, 10.5)
    with c.group():
        c.card(px, c.y, pw, c.h)
        c.tb(px + PAD, c.y + PAD, tw, th, [P(t, b, False, c.t.text_rgb, 8, True) for t in items], anchor=MSO_ANCHOR.MIDDLE)
    return True


def _stat_tiles(c: _Canvas, data) -> bool:
    tiles = (data.get("tiles") or [])[:4]
    if not tiles:
        return False
    n = len(tiles)
    cols = n if n <= 3 else (4 if c.w > 11 else 2)
    cells, ch_full = c.cells(n, cols)
    names = pick_icons(" ".join(str(t.get("label", "")) for t in tiles), n + 2)
    tw = min(w for _, _, w, _ in cells) - 2 * PAD
    th = ch_full - 2 * PAD - ICON_D - 0.12
    def make(vsz, t):
        val = str(t.get("value", ""))
        det = str(t.get("detail", ""))
        lbl = str(t.get("label", ""))
        ps = [P(val, vsz, True, c.t.primary_rgb), P(lbl.upper(), max(12, vsz * 0.34), True, _readable_accent(c.t), 2)]
        if det and not _same(det, lbl):
            ps.append(P(det, max(12, vsz * 0.3), False, c.t.text_rgb, 4))
        return ps
    vsz = min(c.fit_paras(lambda s, t=t: make(s, t), tw, th, 60, 22) for t in tiles)
    need = max(L.paras_height([(p.text, p.size, p.bold, p.before) for p in make(vsz, t)], tw) for t in tiles) + 2 * PAD + ICON_D + 0.12
    if need < ch_full:
        cells, ch_full = c.cells(n, cols, max_ch=max(need * 1.25, ch_full * 0.6))
    for i, (t, (x, y, cw, ch)) in enumerate(zip(tiles, cells)):
        with c.group():
            c.card(x, y, cw, ch)
            c.icon_circle(x + PAD, y + PAD, ICON_D, names[i % len(names)], fill=_blend(c.t.accent_rgb, 0.7), color=_readable_accent(c.t))
            c.tb(x + PAD, y + PAD + ICON_D + 0.12, cw - 2 * PAD, ch - 2 * PAD - ICON_D - 0.12, make(vsz, t), anchor=MSO_ANCHOR.BOTTOM)
    return True


def _cause_effect(c: _Canvas, data) -> bool:
    nodes = (data.get("nodes") or [])[:6]
    if len(nodes) < 3:
        return False
    items = [(str(nd.get("label", "")), str(nd.get("detail", ""))) for nd in nodes]
    return _flow(c, items, False, "↺  and the effect feeds back into the cause" if data.get("loop") else "")


def _stack_height(parts: Sequence[float], gap: float) -> float:
    return sum(parts) + gap * (len(parts) - 1)


_SUB_CHARS = "₀₁₂₃₄₅₆₇₈₉"


def _subscript_runs(shape) -> None:
    """Chemical / algebraic subscripts as real subscript runs (baseline shift) — Unicode subscript glyphs are missing from many fonts."""
    import copy

    for p in shape.text_frame.paragraphs:
        for r in list(p.runs):
            t = r.text
            if not any(ch in _SUB_CHARS for ch in t):
                continue
            segs, cur, cur_sub = [], "", None
            for ch in t:
                is_sub = ch in _SUB_CHARS
                if cur_sub is None or is_sub == cur_sub:
                    cur, cur_sub = cur + ch, is_sub
                else:
                    segs.append((cur, cur_sub))
                    cur, cur_sub = ch, is_sub
            segs.append((cur, cur_sub))
            prev = r._r
            for k, (txt, sub) in enumerate(segs):
                node = prev if k == 0 else copy.deepcopy(r._r)
                if k > 0:
                    prev.addnext(node)
                text = txt.translate(str.maketrans(_SUB_CHARS, "0123456789")) if sub else txt
                for child in node.iterchildren():
                    if child.tag.endswith("}t"):
                        child.text = text
                rpr = node.get_or_add_rPr()
                if sub:
                    rpr.set("baseline", "-25000")
                elif "baseline" in rpr.attrib:
                    del rpr.attrib["baseline"]
                prev = node


_CHEM = re.compile(r"[A-Z][a-z]?\d*\s*(?:\+|⇌|→|<->|->)")


def put_equation(shape, text: str, size: float, color: RGBColor, align: str = "ctr") -> bool:
    """Replace a shape's text by a native PowerPoint equation (Office Math, as in Word); False leaves it untouched.
    Chemical equations keep their text form (their subscripts are handled by _subscript_runs)."""
    from learnova.visuals.mathfmt import looks_like_equation, pretty, to_omml_paragraph

    if not looks_like_equation(text) or _CHEM.search(text):
        return False
    try:
        para = parse_xml(to_omml_paragraph(text, size, _hex(color), fallback=pretty(text), align=align))
    except Exception:
        return False
    body = shape.text_frame._txBody
    for p_ in body.findall(f"{{{_A}}}p"):
        body.remove(p_)
    body.append(para)
    return True


def _formula(c: _Canvas, data) -> bool:
    f = str(data.get("formula", "")).strip()
    if not f:
        return False
    notes = [str(x) for x in data.get("notes") or []][:4]
    term, statement = str(data.get("term", "")), str(data.get("statement", ""))
    label = str(data.get("label", "")).strip()
    gap = 0.2
    term_h = 0.4 if term else 0.0
    label_h = 0.3 if label else 0.0
    st_sz = c.fit([statement], c.w, 1.1, 24, 14) if statement else 0
    st_h = (L.block_height([statement], c.w, st_sz) + 0.05) if statement else 0.0
    notes_h_avail = max(0.8, c.h - term_h - st_h - 1.5 - gap * 3) if notes else 0.0
    n_sz = c.fit(notes, c.w - 0.3, notes_h_avail, 18, 11, spacing=8) if notes else 0
    n_h = (L.block_height(notes, c.w - 0.3, n_sz, 8) + 0.1) if notes else 0.0
    fixed = [h for h in (term_h, st_h, label_h, n_h) if h]
    box_h = min(2.4, max(1.4, c.h - sum(fixed) - gap * len(fixed)))
    parts = [h for h in (term_h, st_h, label_h, box_h, n_h) if h]
    y = c.y + max(0.0, (c.h - _stack_height(parts, gap)) / 2)
    if term:
        with c.group():
            c.tb(c.x, y, c.w, term_h, [P(term.upper(), 15, True, _readable_accent(c.t))], MSO_ANCHOR.MIDDLE)
        y += term_h + gap
    if statement:
        with c.group():
            c.tb(c.x, y, c.w, st_h, [P(statement, st_sz, False, c.t.text_rgb)], MSO_ANCHOR.TOP)
        y += st_h + gap
    if label:
        with c.group():
            c.tb(c.x, y, c.w, label_h, [P(label.upper(), 12, True, c.t.subtext_rgb)], MSO_ANCHOR.BOTTOM)
        y += label_h + gap * 0.5
    with c.group():
        box = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, c.x, y, c.w, box_h, c.t.primary_rgb, radius=0.25 / max(0.5, box_h) * 0.6)
        fsz = c.fit([f], c.w - 2 * PAD - 0.4, box_h - 0.3, 48, 20, bold=True)
        c.tb(0, 0, 0, 0, [P(f, fsz, True, _on(c.t.primary_rgb))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=box, inset=(PAD, 0.1, PAD, 0.1))
        if not put_equation(box, f, min(fsz, 40), _on(c.t.primary_rgb)):
            _subscript_runs(box)
    y += box_h + gap
    if notes:
        with c.group():
            c.tb(c.x, y, c.w, n_h, [P(t, n_sz, False, c.t.text_rgb, 8, True) for t in notes])
    return True


def _definition(c: _Canvas, data) -> bool:
    term, meaning = str(data.get("term", "")), str(data.get("definition", ""))
    if not (term and meaning):
        return False
    notes = [str(x) for x in data.get("notes") or []][:4]
    gap = GUT
    note_cols = notes
    # definition card
    tw = c.w - 2 * PAD - 0.2
    box_max = c.h * (0.62 if notes else 1.0)
    def make(b): return [P(term, b * 1.7, True, c.t.primary_rgb), P(meaning, b, False, c.t.text_rgb, 8)]
    b = c.fit_paras(make, tw, box_max - 2 * PAD, 26, 12)
    need = L.paras_height([(p.text, p.size, p.bold, p.before) for p in make(b)], tw) + 2 * PAD
    box_h = min(box_max, max(need * 1.15, 2.0)) if notes else min(box_max, max(need * 1.4, c.h * 0.7))
    # notes row: small cards under the definition
    cols = min(len(notes), 3) if notes else 0
    note_h = 0.0
    nsz = 0.0
    if notes:
        note_h = min(c.h - box_h - gap, 2.0)
        ncw = (c.w - gap * (cols - 1)) / cols if len(notes) <= 3 else (c.w - gap) / 2
        nsz = c.fit_paras(lambda s: [P(t, s, False, c.t.text_rgb) for t in notes[:1]] and [P(max(notes, key=len), s, False)], ncw - 2 * PAD, (note_h if len(notes) <= 3 else (note_h - gap) / 2) - 2 * PAD, 18, 10.5)
    total = box_h + (gap + note_h if notes else 0)
    y = c.y + max(0.0, (c.h - total) / 2)
    with c.group():
        c.card(c.x, y, c.w, box_h, fill=_blend(c.t.primary_rgb, 0.9), line=c.t.primary_rgb, radius=0.1)
        c.accent_bar(c.x, y + PAD, box_h - 2 * PAD, 0.1)
        c.tb(c.x + PAD + 0.2, y + PAD, tw, box_h - 2 * PAD, make(b), anchor=MSO_ANCHOR.MIDDLE)
    if notes:
        y2 = y + box_h + gap
        if len(notes) <= 3:
            cells = [(c.x + k * ((c.w - gap * (cols - 1)) / cols + gap), y2, (c.w - gap * (cols - 1)) / cols, note_h) for k in range(len(notes))]
        else:
            cells, _ = c.cells(len(notes), 2, area=(c.x, y2, c.w, note_h))
        for t, (x, yy, cw, ch) in zip(notes, cells):
            with c.group():
                c.card(x, yy, cw, ch)
                c.tb(x + PAD, yy + PAD, cw - 2 * PAD, ch - 2 * PAD, [P(t, nsz, False, c.t.text_rgb)], anchor=MSO_ANCHOR.MIDDLE)
    return True


def _quote(c: _Canvas, data) -> bool:
    t = str(data.get("text", ""))
    if not t:
        return False
    attr = str(data.get("attribution") or "")
    tw, th = c.w - 2 * PAD - 0.5, c.h - 2 * PAD - (0.6 if attr else 0)
    b = c.fit_paras(lambda s: [P("“" + t + "”", s, True, c.t.primary_rgb)], tw, th, 28, 14)
    need = L.paras_height([("“" + t + "”", b, True, 0)], tw) + (0.6 if attr else 0) + 2 * PAD
    bh = min(c.h, max(need * 1.25, c.h * 0.55))
    y = c.y + (c.h - bh) / 2
    with c.group():
        c.card(c.x, y, c.w, bh, fill=_blend(c.t.primary_rgb, 0.9), line=_blend(c.t.primary_rgb, 0.6))
        c.accent_bar(c.x, y + PAD, bh - 2 * PAD, 0.12)
        paras = [P("“" + t + "”", b, True, c.t.primary_rgb)]
        if attr:
            paras.append(P("— " + attr, max(13, b * 0.5), False, c.t.text_rgb, 10))
        c.tb(c.x + PAD + 0.3, y + PAD, tw, bh - 2 * PAD, paras, anchor=MSO_ANCHOR.MIDDLE)
    return True


def _funnel(c: _Canvas, data) -> bool:
    st = (data.get("stages") or [])[:6]
    if len(st) < 3:
        return False
    n = len(st)
    gap = GUT * 0.5
    rh_full = (c.h - gap * (n - 1)) / n
    rh = min(rh_full, 1.1)
    oy = (c.h - (rh * n + gap * (n - 1))) / 2
    narrow = 0.4
    labels = [f"{s_.get('value', '')}   {s_.get('label', '')}".strip() for s_ in st]
    min_w = c.w * (1 - narrow * (n - 1) / max(1, n - 1))
    b = min(c.fit([lb], c.w * (1 - narrow) - 2 * PAD, rh - 0.14, 24, 11, bold=True) for lb in labels)
    for i, (s_, lb) in enumerate(zip(st, labels)):
        w = c.w * (1.0 - i * (narrow / max(1, n - 1)))
        fill = _blend(c.t.primary_rgb, 0.05 + 0.5 * i / max(1, n - 1))
        y = c.y + oy + i * (rh + gap)
        with c.group():
            shp = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, c.x + (c.w - w) / 2, y, w, rh, fill, radius=0.15)
            c.tb(0, 0, 0, 0, [P(lb, b, True, _on(fill))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=shp, inset=(PAD, 0.04, PAD, 0.04))
    return True


def _chart(c: _Canvas, data, kind: str) -> bool:
    pts = []
    for p in data.get("points") or data.get("series") or data.get("items") or []:
        try:
            pts.append((str(p.get("label") or p.get("x") or ""), float(p.get("value") if p.get("value") is not None else p.get("y"))))
        except (AttributeError, TypeError, ValueError):
            continue
    pts = pts[:10]
    if len(pts) < 2:
        return False
    cd = CategoryChartData()
    cd.categories = [l for l, _ in pts]
    cd.add_series("Value", [v for _, v in pts])
    ctype = {"pie": XL_CHART_TYPE.PIE, "line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.BAR_CLUSTERED}[kind]
    with c.group():
        c.card(c.x, c.y, c.w, c.h)
        inset = 0.15
        gf = c.s.shapes.add_chart(ctype, Inches(c.x + inset), Inches(c.y + inset), Inches(c.w - 2 * inset), Inches(c.h - 2 * inset), cd)
        ch = gf.chart
        ch.has_legend = kind == "pie"
        ch.has_title = False                      # the slide title says it; "Value" above the plot is noise
        ch.font.size = Pt(14)
        ch.font.color.rgb = c.t.text_rgb
        if kind == "pie":
            ch.legend.position = XL_LEGEND_POSITION.RIGHT
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(15)
        plot = ch.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(14)
        plot.data_labels.font.bold = True
        fmt = data.get("fmt") or {}
        if fmt.get("prefix") or fmt.get("suffix"):      # "$5.6M", "62%": the figures as the source wrote them
            decimals = "0.0" if any(abs(v - round(v)) > 1e-9 for _, v in pts) else "0"
            plot.data_labels.number_format = f'"{fmt.get("prefix", "")}"#,##{decimals}"{fmt.get("suffix", "")}"'
            plot.data_labels.number_format_is_linked = False
        if kind == "pie":
            for i in range(len(pts)):
                pt = plot.series[0].points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _blend(c.t.primary_rgb, 0.0 if i == 0 else min(0.75, 0.18 * i))
        else:
            ser = plot.series[0]
            if kind == "bar":
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = c.t.primary_rgb
                plot.gap_width = 60
                ch.category_axis.reverse_order = True
            else:
                ser.format.line.color.rgb = c.t.primary_rgb
                ser.format.line.width = Pt(3.5)
                ser.smooth = False
            ch.category_axis.tick_labels.font.size = Pt(14)
            ch.value_axis.tick_labels.font.size = Pt(13)
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = _blend(c.t.primary_rgb, 0.85)
    return True


def _function_plot(c: _Canvas, data) -> bool:
    from learnova.rendering.family_blocks import _eval_expr

    expr = str(data.get("expr", "")).strip()
    a, b = (data.get("domain") or [-5, 5])[:2]
    xs = [a + (b - a) * i / 80 for i in range(81)]
    ys = [(x, _eval_expr(expr, x)) for x in xs]
    pts = [(x, y) for x, y in ys if y is not None and abs(y) < 1e6]
    if len(pts) < 6:
        return False
    ymin, ymax = min(y for _, y in pts), max(y for _, y in pts)
    if ymax == ymin:
        ymax += 1
    px, py, pw, ph = c.x + PAD, c.y + PAD, c.w - 2 * PAD, c.h - 2 * PAD - 0.5
    with c.group():
        c.card(c.x, c.y, c.w, c.h)
        c.shape(MSO_SHAPE.RECTANGLE, px, py, pw, ph, _blend(c.t.primary_rgb, 0.95), _blend(c.t.primary_rgb, 0.6))
        if ymin < 0 < ymax:
            yy = py + ph * (1 - (0 - ymin) / (ymax - ymin))
            c.shape(MSO_SHAPE.RECTANGLE, px, yy - 0.01, pw, 0.02, c.t.text_rgb)
        if a < 0 < b:
            xx = px + pw * (0 - a) / (b - a)
            c.shape(MSO_SHAPE.RECTANGLE, xx - 0.01, py, 0.02, ph, c.t.text_rgb)
        coords = [(Inches(px + pw * (x - a) / (b - a)), Inches(py + ph * (1 - (y - ymin) / (ymax - ymin)))) for x, y in pts]
        fb = c.s.shapes.build_freeform(Emu(int(coords[0][0])), Emu(int(coords[0][1])), scale=1.0)
        fb.add_line_segments([(Emu(int(x)), Emu(int(y))) for x, y in coords[1:]], close=False)
        curve = fb.convert_to_shape()
        curve.fill.background()
        curve.line.color.rgb = c.t.primary_rgb
        curve.line.width = Pt(3.5)
        c.tb(px, py + ph + 0.05, pw, 0.4, [P(f"y = {expr}", 16, True, c.t.primary_rgb)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
    return True


def _mind_map(c: _Canvas, data) -> bool:
    center = str(data.get("center") or data.get("central") or "")
    br = [str(b) for b in (data.get("branches") or data.get("items") or []) if str(b).strip()][:8]
    if not center or len(br) < 3:
        return False
    cx, cy = c.x + c.w / 2, c.y + c.h / 2
    bw = min(3.3, c.w / 3.4)
    bh = min(1.1, c.h / 3.6)
    hub_w, hub_h = min(3.0, c.w * 0.26), min(1.5, c.h * 0.26)
    rx, ry = c.w / 2 - bw / 2, c.h / 2 - bh / 2
    pos = []
    for i in range(len(br)):
        ang = 2 * math.pi * i / len(br) - math.pi / 2
        pos.append((cx + rx * math.cos(ang), cy + ry * math.sin(ang)))
    b = min(c.fit([t], bw - 2 * PAD * 0.7, bh - 0.16, 16, 10) for t in br)
    with c.group():
        for x, y in pos:   # spokes sit behind everything else
            ln = c.s.shapes.add_connector(1, Inches(cx), Inches(cy), Inches(x), Inches(y))
            ln.line.color.rgb = _blend(c.t.primary_rgb, 0.5)
            ln.line.width = Pt(2)
        hub = c.shape(MSO_SHAPE.OVAL, cx - hub_w / 2, cy - hub_h / 2, hub_w, hub_h, c.t.primary_rgb)
        hsz = c.fit([center], hub_w * 0.75, hub_h * 0.7, 22, 11, bold=True)
        c.tb(0, 0, 0, 0, [P(center, hsz, True, _on(c.t.primary_rgb))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=hub, inset=(0.1, 0.05, 0.1, 0.05))
    for t, (x, y) in zip(br, pos):
        with c.group():
            c.card(x - bw / 2, y - bh / 2, bw, bh)
            c.tb(x - bw / 2 + PAD * 0.7, y - bh / 2 + 0.08, bw - PAD * 1.4, bh - 0.16, [P(t, b, False, c.t.text_rgb)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
    return True


_REL_RGB = {
    "contrast": RGBColor(0xC2, 0x41, 0x0C), "concession": RGBColor(0xC2, 0x41, 0x0C), "cause": RGBColor(0x1D, 0x4E, 0xD8),
    "result": RGBColor(0x1D, 0x4E, 0xD8), "example": RGBColor(0x0F, 0x76, 0x6E), "addition": RGBColor(0x0F, 0x76, 0x6E),
    "sequence": RGBColor(0x6D, 0x28, 0xD9), "purpose": RGBColor(0x6D, 0x28, 0xD9), "elaboration": RGBColor(0x64, 0x74, 0x8B),
}
_REL_TAG = {"contrast": "but", "concession": "although", "cause": "because", "result": "so", "example": "e.g.",
            "addition": "also", "sequence": "then", "purpose": "in order to"}


def _prose(c: _Canvas, data) -> bool:
    """Plain text, no boxes: an emphasised lead sentence, then the rest as prose with relations marked
    (coloured connectives, a small relation tag, examples indented)."""
    from learnova.visuals.discourse import cue_spans

    lead = str(data.get("lead", "")).strip()
    body = [(str(b.get("text", "")).strip(), str(b.get("rel", "none"))) for b in (data.get("body") or []) if str(b.get("text", "")).strip()]
    if not lead:
        return False
    bar_w = 0.09
    facts = [f for f in (data.get("facts") or []) if f.get("value")][:3]
    rail_w = 3.1 if facts else 0.0
    text_w = c.w - (rail_w + GUT if facts else 0.0)
    lx = c.x + bar_w + PAD * 0.9
    lw = (text_w - (lx - c.x)) * 0.93   # conservative: real fonts wrap earlier than the estimate
    gap = 0.14

    def tagged(text, rel):
        tag = _REL_TAG.get(rel, "")
        return (tag.upper() + "   " if tag else "") + text

    def _pw(i):
        """Width used to estimate a paragraph: the bold lead face (Poppins) runs wider than the estimate, so it gets extra margin."""
        return lw * 0.88 if i == 0 else lw - (0.3 if body[i - 1][1] in {"example", "elaboration"} else 0)

    lead_k = 1.38 if len(lead.split()) <= 14 else 1.24 if len(lead.split()) <= 24 else 1.12     # a long lead sentence must not shout

    def make(b):
        ps = [(lead, b * lead_k, True, 0)] + [(tagged(t, r), b, False, 0) for t, r in body]
        return ps

    b = 24.0
    while b > 11.0:
        ps = make(b)
        total = sum(L.paras_height([p], _pw(i)) for i, p in enumerate(ps)) + gap * (len(ps) + 0.5)
        if total <= c.h * 0.96 and L.longest_word_fits([p[0] for p in ps], lw, b, False):
            break
        b -= 0.5
    ps = make(b)
    heights = [L.paras_height([p], _pw(i)) for i, p in enumerate(ps)]
    total = sum(heights) + gap * (len(ps) + 0.5)
    y = c.y + max(0.0, (c.h - total) / 2)
    with c.group():
        c.accent_bar(c.x, y + 0.04, heights[0] - 0.02, bar_w)
        c.tb(lx, y, lw, heights[0], [P(lead, b * lead_k, True, c.t.primary_rgb)], anchor=MSO_ANCHOR.TOP)
    y += heights[0] + gap * 1.5
    for (text, rel), h in zip(body, heights[1:]):
        indent = 0.3 if rel in {"example", "elaboration"} else 0.0
        with c.group():
            box = c.s.shapes.add_textbox(Inches(lx + indent), Inches(y), Inches(lw - indent), Inches(h))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            tag = _REL_TAG.get(rel, "")
            if tag:
                r0 = p.add_run()
                r0.text = tag.upper() + "   "
                r0.font.size = Pt(round(b * 0.62, 1))
                r0.font.bold = True
                r0.font.color.rgb = _REL_RGB.get(rel, c.t.primary_rgb)
            pos = 0
            for a, z, name in cue_spans(text) + [(len(text), len(text), "")]:
                if a > pos:
                    r1 = p.add_run()
                    r1.text = text[pos:a]
                    r1.font.size = Pt(round(b, 1))
                    r1.font.color.rgb = c.t.text_rgb
                if z > a:
                    r2 = p.add_run()
                    r2.text = text[a:z]
                    r2.font.size = Pt(round(b, 1))
                    r2.font.bold = True
                    r2.font.color.rgb = _REL_RGB.get(name, c.t.primary_rgb)
                pos = max(pos, z)
        y += h + gap
    if facts:
        _facts_rail(c, facts, c.x + c.w - rail_w, rail_w)
    return True


def _outline(c: _Canvas, data) -> bool:
    """Nested bullets as grouped cards: heading + its points per card, a lead line above, childless siblings as chips below."""
    groups = [g for g in (data.get("groups") or []) if g.get("heading")][:6]
    chips = [str(x) for x in (data.get("chips") or []) if str(x).strip()][:12]
    lead = str(data.get("lead", "")).strip()
    more = [str(x) for x in (data.get("lead_more") or []) if str(x).strip()]
    if not groups:
        return False
    gap = GUT
    lead_h = 0.0
    lead_sz = 0.0
    if lead:
        lead_sz = c.fit([lead], c.w, 0.9, 24, 15, bold=True)
        lead_h = L.block_height([lead], c.w, lead_sz, bold=True) + 0.08
        if more:
            lead_h += L.block_height(more, c.w - 0.3, 14, 4) + 0.1
    chip_rows = 0
    chip_h = 0.42
    if chips:
        per_row = max(1, int(c.w // 2.3))
        chip_rows = math.ceil(len(chips) / per_row)
    chips_total = (chip_rows * chip_h + (chip_rows - 1) * 0.12 + gap) if chips else 0.0
    grid_h = c.h - (lead_h + gap if lead else 0.0) - chips_total
    n = len(groups)
    cols = n if n <= 3 else 3 if n in (5, 6) else 2
    rows = math.ceil(n / cols)
    cw = (c.w - gap * (cols - 1)) / cols
    chh = (grid_h - gap * (rows - 1)) / rows

    def make(b, g):
        ps = [P(str(g["heading"]), b * 1.25, True, c.t.primary_rgb)]
        for it in g.get("items") or []:
            sub = str(it).startswith("– ")
            ps.append(P(str(it)[2:] if sub else str(it), b * (0.9 if sub else 1.0), False, c.t.subtext_rgb if sub else c.t.text_rgb, 3 if sub else 7, True, 1 if sub else 0))
        return ps

    fits = [c.fit_paras(lambda b, g_=g: make(b, g_), cw - 2 * PAD, chh - 2 * PAD, 20, 10) for g in groups]
    b = _shared(fits)

    def fits_at_b(g_) -> bool:
        return L.paras_height([(q.text, q.size, q.bold, q.before) for q in make(b, g_)], cw - 2 * PAD) <= chh - 1.6 * PAD

    trimmed = []
    for g in groups:                                       # a card that cannot hold all its points at a readable size drops the last ones
        items = list(g.get("items") or [])
        dropped = 0
        while len(items) > 1 and not fits_at_b({**g, "items": items + ([f"+{dropped} more"] if dropped else [])}):
            items.pop()
            dropped += 1
        trimmed.append({**g, "items": items + ([f"+{dropped} more"] if dropped else [])})
    groups = trimmed
    y = c.y
    if lead:
        with c.group():
            c.accent_bar(c.x, y + 0.03, lead_h - 0.1, 0.08)
            c.tb(c.x + 0.25, y, c.w - 0.25, lead_h, [P(lead, lead_sz, True, c.t.primary_rgb)] + [P(m, 14, False, c.t.text_rgb, 4) for m in more], MSO_ANCHOR.MIDDLE)
        y += lead_h + gap
    for i, g in enumerate(groups):
        r, k = divmod(i, cols)
        x = c.x + k * (cw + gap)
        yy = y + r * (chh + gap)
        with c.group():
            c.card(x, yy, cw, chh)
            c.accent_bar(x, yy + 0.18, 0.55, 0.07)
            c.tb(x + PAD, yy + PAD * 0.8, cw - 2 * PAD, chh - 1.6 * PAD, make(b, g), anchor=MSO_ANCHOR.TOP, bullet_color=c.t.accent_rgb)
    if chips:
        cy = y + rows * chh + (rows - 1) * gap + gap
        per_row = max(1, int(c.w // 2.3))
        with c.group():
            for k, t in enumerate(chips):
                r, kk = divmod(k, per_row)
                w = (c.w - 0.14 * (per_row - 1)) / per_row
                x = c.x + kk * (w + 0.14)
                pill = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, cy + r * (chip_h + 0.12), w, chip_h, _blend(c.t.primary_rgb, 0.88), _blend(c.t.primary_rgb, 0.6), radius=0.5)
                c.tb(0, 0, 0, 0, [P(t, c.fit([t], w - 0.3, chip_h - 0.08, 15, 10, bold=True), True, c.t.primary_rgb)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=pill, inset=(0.1, 0.02, 0.1, 0.02))
    return True


def _facts_rail(c: _Canvas, facts, x: float, w: float) -> None:
    """"At a glance": the slide's own figures as stacked cards beside the text — a visual anchor made from the content."""
    n = len(facts)
    head_h = 0.34
    gap = 0.18
    ch = min(1.6, (c.h - head_h - gap * n) / n)
    total = head_h + gap + n * ch + (n - 1) * gap
    y = c.y + max(0.0, (c.h - total) / 2)
    with c.group():
        c.tb(x, y, w, head_h, [P("AT A GLANCE", 11, True, c.t.subtext_rgb)], MSO_ANCHOR.MIDDLE)
    y += head_h + gap
    for f in facts:
        with c.group():
            c.card(x, y, w, ch, fill=_blend(c.t.primary_rgb, 0.92), line=_blend(c.t.primary_rgb, 0.7))
            c.accent_bar(x + 0.12, y + 0.18, ch - 0.36, 0.07)
            vsz = c.fit([str(f["value"])], w - 0.55, ch * 0.5, 36, 16, bold=True)
            lsz = c.fit([str(f.get("label", ""))], w - 0.55, ch * 0.34, 14, 10)
            c.tb(x + 0.3, y + 0.08, w - 0.45, ch - 0.16,
                 [P(str(f["value"]), vsz, True, c.t.primary_rgb), P(str(f.get("label", "")), lsz, False, c.t.subtext_rgb, 2)], MSO_ANCHOR.MIDDLE)
        y += ch + gap


def _sections(c: _Canvas, data) -> bool:
    """Sibling subsections side by side (no boxes): accent rule, heading, lead sentence, text — one shared type size."""
    secs = (data.get("sections") or [])[:3]
    if len(secs) < 2:
        return False
    n = len(secs)
    cw = (c.w - GUT * 1.4 * (n - 1)) / n
    gap = 0.14

    def paras(sc, b):
        out = [P(str(sc.get("title", "")), b * 1.12, True, c.t.primary_rgb),
               P(str(sc.get("lead", "")), b * 1.04, True, c.t.text_rgb, 8)]
        body = " ".join(str(x) for x in sc.get("body", []))
        if body:
            out.append(P(body, b, False, c.t.text_rgb, 8))
        return out

    b = 22.0
    while b > 9.5:
        if all(L.paras_height([(p.text, p.size, p.bold, p.before) for p in paras(sc, b)], cw) + 0.2 <= c.h * 0.97 for sc in secs):
            break
        b -= 0.5
    hs = [L.paras_height([(p.text, p.size, p.bold, p.before) for p in paras(sc, b)], cw) for sc in secs]
    top = c.y + max(0.0, (c.h - (max(hs) + 0.2)) / 2)
    for i, sc in enumerate(secs):
        x = c.x + i * (cw + GUT * 1.4)
        with c.group():
            c.shape(MSO_SHAPE.RECTANGLE, x, top, cw, 0.07, c.t.accent_rgb)
            c.tb(x, top + 0.16, cw, hs[i] + 0.05, paras(sc, b), anchor=MSO_ANCHOR.TOP)
    return True


# ── generic entry points used by ppt_builder ─────────────────────────────────
def draw_cards(slide, theme, band, items: List[Tuple[str, str]], kind: str = "grid", numbered: bool = False) -> bool:
    """Card grid or arrow-linked process from ``(heading, body)`` pairs (heading may be empty)."""
    items = [(h, b) for h, b in items if (h or b).strip()]
    if not items:
        return False
    c = _Canvas(slide, theme, band)
    if kind == "process" and len(items) >= 2:
        return _flow(c, items)
    _grid(c, items, numbered=numbered, icons=not numbered)
    return True


def draw_bullets(slide, theme, band, bullets: List[str]) -> bool:
    """Bullets as full-width rows (accent bar + hanging-indent bullet text), one animation group per bullet."""
    bullets = [b for b in bullets if b.strip()][:9]
    if not bullets:
        return False
    c = _Canvas(slide, theme, band)
    n = len(bullets)
    gap = GUT * 0.6
    rh_full = (c.h - gap * (n - 1)) / n
    tw = c.w - 2 * PAD - 0.3
    b = c.fit_paras(lambda s: [P(max(bullets, key=len), s, False)], tw - 0.3, rh_full - 0.16, 26, 11)
    need = max(L.paras_height([(t, b, False, 0)], tw - 0.3) for t in bullets) + 0.3
    rh = min(rh_full, max(need * 1.15, 0.75, rh_full * 0.6), 1.6)
    oy = (c.h - (rh * n + gap * (n - 1))) / 2
    for i, t in enumerate(bullets):
        y = c.y + oy + i * (rh + gap)
        with c.group():
            c.card(c.x, y, c.w, rh, fill=_blend(theme.primary_rgb, 0.93), line=_blend(theme.primary_rgb, 0.8))
            c.accent_bar(c.x, y + 0.14, rh - 0.28, 0.09)
            c.tb(c.x + PAD + 0.12, y + 0.08, tw, rh - 0.16, [P(t, b, False, theme.text_rgb, 0, True)], anchor=MSO_ANCHOR.MIDDLE)
    return True


def draw_metric(slide, theme, band, value: str, label: str, desc: str) -> bool:
    c = _Canvas(slide, theme, band)
    h = min(c.h, 4.8)
    y = c.y + (c.h - h) / 2
    with c.group():
        box = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, c.x, y, c.w, h, theme.primary_rgb, theme.accent_rgb, 3, radius=0.12 / max(0.5, h) * 2)
        tw, th = c.w - 2 * PAD - 0.6, h - 2 * PAD
        vc = theme.accent_rgb if _lum(theme.accent_rgb, theme.primary_rgb) else _on(theme.primary_rgb)
        def make(v): return [P(value, v, True, vc),
                             P(label.upper(), max(14, v * 0.3), True, _on(theme.primary_rgb), 8),
                             P(desc, max(13, v * 0.22), False, _on(theme.primary_rgb), 8)]
        v = c.fit_paras(make, tw, th, 96, 36)
        c.tb(0, 0, 0, 0, make(v), MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=box, inset=(PAD + 0.3, PAD, PAD + 0.3, PAD))
    return True


def _lum(a: RGBColor, bg: RGBColor) -> bool:
    def l(rgb):
        ch = []
        for v in (rgb[0], rgb[1], rgb[2]):
            x = v / 255.0
            ch.append(x / 12.92 if x <= 0.04045 else ((x + 0.055) / 1.055) ** 2.4)
        return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]
    hi, lo = sorted((l(a), l(bg)), reverse=True)
    return (hi + 0.05) / (lo + 0.05) >= 3.0


_KW = re.compile(r"^(step\s*\d+|input|output|begin|end|for each|foreach|for|while|repeat|until|do|if|else|elif|then|return|initiali[sz]e|init|"
                 r"set|update|compute|calculate|let|loop|break|continue|print|def|function|procedure|call|stop|exit)\b:?", re.I)


def _algorithm(c: _Canvas, data) -> bool:
    """Pseudocode or a derivation in ONE panel: numbered lines, indentation kept, keywords bold (code) or centred maths (math)."""
    lines = [(str(x.get("text", "")), int(x.get("indent", 0))) for x in (data.get("lines") or []) if str(x.get("text", "")).strip()]
    if len(lines) < 2:
        return False
    math = data.get("variant") == "math" or str(data.get("label", "")).lower().startswith("equation")
    head_h = 0.42
    with c.group():
        c.card(c.x, c.y, c.w, c.h, fill=_blend(c.t.primary_rgb, 0.96), line=_blend(c.t.primary_rgb, 0.7), radius=0.14)
        strip = c.shape(MSO_SHAPE.RECTANGLE, c.x, c.y, c.w, head_h, c.t.primary_rgb)
        c.tb(0, 0, 0, 0, [P(str(data.get("label") or ("Equations" if math else "Algorithm")).upper(), 13, True, _on(c.t.primary_rgb))],
             MSO_ANCHOR.MIDDLE, shape=strip, inset=(PAD, 0, PAD, 0))
    num_w = 0.0 if math else 0.45
    tx, tw = c.x + PAD + num_w, c.w - 2 * PAD - num_w
    avail = c.h - head_h - 0.3
    size = 22.0 if math else 18.0
    while size > 11:
        need = sum(L.estimate_lines(t, tw - 0.32 * d, size) * size * 1.32 / 72 for t, d in lines)
        if need <= avail and L.longest_word_fits([t for t, _ in lines], tw, size):
            break
        size -= 0.5
    row = [L.estimate_lines(t, tw - 0.32 * d, size) * size * 1.32 / 72 for t, d in lines]
    y = c.y + head_h + 0.15 + max(0.0, (avail - sum(row)) / 2) * (0.5 if math else 0.0)
    font = "Cambria Math" if math else "Consolas"
    for i, ((t, d), h) in enumerate(zip(lines, row)):
        with c.group():
            if not math:
                nb = c.s.shapes.add_textbox(Inches(c.x + PAD * 0.6), Inches(y), Inches(num_w), Inches(h))
                r = nb.text_frame.paragraphs[0].add_run()
                r.text, r.font.size, r.font.color.rgb = str(i + 1), Pt(size * 0.75), _blend(c.t.text_rgb, 0.5)
            box = c.s.shapes.add_textbox(Inches(tx + 0.32 * d), Inches(y), Inches(tw - 0.32 * d), Inches(h))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p_ = tf.paragraphs[0]
            if math:
                p_.alignment = PP_ALIGN.CENTER
            if math and put_equation(box, t, size, c.t.text_rgb):
                y += h
                continue
            m = None if math else _KW.match(t)
            parts = [(t[:m.end()], True), (t[m.end():], False)] if m else [(t, False)]
            for txt, kw in parts:
                if not txt:
                    continue
                r = p_.add_run()
                r.text = _superscripts(txt) if math else txt
                r.font.size, r.font.name = Pt(size), font
                r.font.bold = kw
                r.font.color.rgb = c.t.primary_rgb if kw else c.t.text_rgb
        y += h
    return True


def _flow_layout(n: int, x: float, y: float, w: float, h: float, lane: float):
    """Boxes for n flowchart nodes: one row up to 7, else two rows; returns [(x, y, w, h)] and the row of each node."""
    rows = 1 if n <= 7 else 2
    per = math.ceil(n / rows)
    gap = 0.42
    bw = min(2.3, (w - gap * (per - 1)) / per)
    bh = min(1.25, (h - lane * rows - 0.25 * (rows - 1)) / rows)
    x0 = x + (w - (bw * per + gap * (per - 1))) / 2
    out, row_of = [], []
    for i in range(n):
        r, k = divmod(i, per)
        out.append((x0 + k * (bw + gap), y + r * (bh + lane + 0.25), bw, bh))
        row_of.append(r)
    return out, row_of, bh


def _line(slide, x1, y1, x2, y2, color, arrow=False, w=1.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(w)
    if arrow:
        ln.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd xmlns:a="{_A}" type="triangle" w="med" len="med"/>'))
    return ln


def _flowchart(c: _Canvas, data) -> bool:
    """Start / process / decision / end shapes joined by arrows; a decision's other branch goes to a side box in the lane
    below, or loops back (Yes / No labelled) to the step it repeats."""
    nodes = [n for n in (data.get("nodes") or []) if isinstance(n, dict) and str(n.get("text", "")).strip()]
    if len(nodes) < 3:
        return False
    lane = 0.95 if any(n.get("kind") == "decision" for n in nodes) else 0.2
    top = 0.35 if len(nodes) > 7 and any(n.get("back") for n in nodes) else 0.05     # room above row 1 for a loop from row 2
    boxes, row_of, bh = _flow_layout(len(nodes), c.x + 0.3, c.y + top, c.w - 0.3, c.h - top - 0.05, lane)
    ink = c.t.text_rgb
    shapes = []
    size = c.fit([str(n["text"]) for n in nodes if n.get("kind") != "decision"], boxes[0][2] - 0.25, bh - 0.2, 16, 10.5)
    dtexts = [str(n["text"]) for n in nodes if n.get("kind") == "decision"]
    dw = (boxes[0][2] + 0.16) * 0.5 - 0.04                # text area of a diamond: the inscribed half-width rectangle
    dsize = c.fit(dtexts, dw, (bh + 0.1) * 0.52, 13, 8.5, bold=True) if dtexts else size
    while dtexts and dsize > 8.5 and not L.longest_word_fits(dtexts, dw, dsize, True):
        dsize -= 0.5                                                   # a diamond's middle is narrow: never split a word there
    for n, (x, y, w, h) in zip(nodes, boxes):
        kind = n.get("kind")
        with c.group():
            if kind in ("start", "end"):
                shp = c.shape(MSO_SHAPE.FLOWCHART_TERMINATOR, x + w * 0.12, y + h * 0.2, w * 0.76, h * 0.6, c.t.primary_rgb)
                c.tb(0, 0, 0, 0, [P(str(n["text"]), size, True, _on(c.t.primary_rgb))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=shp)
            elif kind == "decision":
                shp = c.shape(MSO_SHAPE.DIAMOND, x - 0.08, y - 0.05, w + 0.16, h + 0.1, _blend(c.t.accent_rgb, 0.75), c.t.accent_rgb, 1.5)
                # a text box over the diamond's inscribed middle (viewers differ on text insets inside a diamond)
                c.tb(x + w / 2 - dw / 2, y + h * 0.2, dw, h * 0.6, [P(str(n["text"]), dsize, True, ink)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
            else:
                shp = c.card(x, y, w, h)
                c.tb(0, 0, 0, 0, [P(str(n["text"]), size, False, ink)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=shp,
                     inset=(0.1, 0.05, 0.1, 0.05))
        shapes.append(shp)
    proc_index = [i for i, n in enumerate(nodes) if n.get("kind") == "process"]
    lab = lambda x, y, t: c.tb(x, y, 0.5, 0.26, [P(t, 11, True, _readable_accent(c.t))], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)   # noqa: E731
    for i in range(len(nodes) - 1):
        (x1, y1, w1, h1), (x2, y2, w2, h2) = boxes[i], boxes[i + 1]
        with c.group():
            if row_of[i] == row_of[i + 1]:
                _line(c.s, x1 + w1, y1 + h1 / 2, x2, y2 + h2 / 2, ink, arrow=True)
            else:                                          # wrap to the next row: down, across to the left, down
                mid = y1 + h1 + lane + 0.12
                _line(c.s, x1 + w1 / 2, y1 + h1, x1 + w1 / 2, mid, ink)
                _line(c.s, x1 + w1 / 2, mid, x2 + w2 / 2, mid, ink)
                _line(c.s, x2 + w2 / 2, mid, x2 + w2 / 2, y2, ink, arrow=True)
            if nodes[i].get("kind") == "decision":
                lab(x1 + w1 - 0.05, y1 + h1 / 2 - 0.32, str(nodes[i].get("cont") or "Yes"))
    for i, n in enumerate(nodes):
        if n.get("kind") != "decision":
            continue
        x, y, w, h = boxes[i]
        low = y + h + lane * 0.55
        side = str(n.get("side") or "No")
        with c.group():
            if n.get("back") and proc_index:               # loop back to the step it repeats
                t = proc_index[min(len(proc_index), max(1, int(n["back"]))) - 1]
                tx, ty, tw, th = boxes[t]
                acc = c.t.accent_rgb
                _line(c.s, x + w / 2, y + h, x + w / 2, low, acc, w=2)
                if row_of[t] == row_of[i]:                 # same row: a U under the row
                    _line(c.s, x + w / 2, low, tx + tw / 2, low, acc, w=2)
                    _line(c.s, tx + tw / 2, low, tx + tw / 2, ty + th, acc, arrow=True, w=2)
                else:                                      # earlier row: round the left edge and come in from above
                    lx, hy = boxes[0][0] - 0.2, ty - 0.18
                    _line(c.s, x + w / 2, low, lx, low, acc, w=2)
                    _line(c.s, lx, low, lx, hy, acc, w=2)
                    _line(c.s, lx, hy, tx + tw / 2, hy, acc, w=2)
                    _line(c.s, tx + tw / 2, hy, tx + tw / 2, ty, acc, arrow=True, w=2)
                lab(x + w / 2 + 0.02, y + h + 0.02, side)
            elif n.get("side_text"):                       # the other case: its own box in the lane below
                sw = min(w * 1.25, 2.6)
                bx = x + w / 2 - sw / 2
                _line(c.s, x + w / 2, y + h, x + w / 2, y + h + 0.22, c.t.accent_rgb, arrow=True, w=2)
                box = c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y + h + 0.24, sw, lane - 0.3, _blend(c.t.accent_rgb, 0.88), c.t.accent_rgb, 1.25)
                c.tb(0, 0, 0, 0, [P(str(n["side_text"]), min(size, 13), False, ink)], MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, shape=box,
                     inset=(0.06, 0.02, 0.06, 0.02))
                lab(x + w / 2 + 0.02, y + h - 0.02, side)
    return True


_NO_PANEL = {"FLOWCHART", "ALGORITHM", "COMPARE_TABLE", "REFERENCES", "SECTIONS", "PROSE", "FORMULA_CALLOUT", "DEFINITION", "QUOTE", "WORKED_EXAMPLE", "STEPPER"}


def explain_panel(explain: str, notes: Sequence[str], band, family: str = ""):
    """(height, font size, label, paragraphs) of the "In plain words" panel under a visual, or None.

    The panel carries the AI's plain explanation and the slide's own sentences that the visual does not show, as readable
    paragraphs (11–14 pt) instead of a one-line strip. It may take at most 40% of the band; sentences that do not fit are
    dropped from the end (they stay in the speaker notes), never shrunk below 11 pt."""
    explain = explain.strip()
    notes = [n.strip() for n in notes if n.strip() and n.strip() not in explain]
    if family in _NO_PANEL and not explain:
        return None
    paras = ([explain] if explain else []) + ([" ".join(notes)] if notes else [])
    if not paras:
        return None
    label = "IN PLAIN WORDS" if explain else "ALSO ON THIS SLIDE"
    width = band.width - 0.08 - PAD * 1.6
    cap = band.height * 0.4
    while paras:
        for size in (14.0, 13.0, 12.0, 11.0):
            h = L.paras_height([(label, size * 0.72, True, 0)] + [(t, size, False, 3) for t in paras], width) * 1.08 + 0.2
            if h <= cap:
                return (max(h, 0.62), size, label, paras)
        last = paras[-1]
        cut = re.split(r"(?<=[.!?])\s+", last)
        if len(cut) > 1:
            paras[-1] = " ".join(cut[:-1])
        else:
            paras.pop()
    return None


def draw_family(slide, sp, theme, band) -> bool:
    family = getattr(sp, "family", "") or ""
    data = getattr(sp, "data", {}) or {}
    if not data:
        return False
    params = getattr(sp, "params", None) or {}
    notes = [str(n) for n in (params.get("notes") or []) if str(n).strip()]
    panel = explain_panel(str(params.get("explain") or ""), notes, band, family)
    note_h = panel[0] if panel else 0.0
    c = _Canvas(slide, theme, band, reserve=(note_h + GUT) if panel else 0.0)
    try:
        if family in {"PROCESS_LINEAR", "PROCESS_CYCLIC"}:
            ok = _process(c, data, family == "PROCESS_CYCLIC")
        elif family == "STEPPER":
            ok = _stepper(c, data)
        elif family == "WORKED_EXAMPLE":
            ok = _worked(c, data)
        elif family == "TIMELINE":
            ok = _timeline(c, data)
        elif family == "COMPARE_TABLE":
            ok = _compare_table(c, data)
        elif family == "COMPARE_VISUAL":
            ok = _pros_cons(c, data)
        elif family == "MATRIX_GRID":
            ok = _matrix(c, data)
        elif family == "REFERENCES":
            ok = _references(c, data)
        elif family == "PROSE":
            ok = _prose(c, data)
        elif family == "FLOWCHART":
            ok = _flowchart(c, data)
        elif family == "ALGORITHM":
            ok = _algorithm(c, {**data, "variant": getattr(sp, "variant", "") or data.get("variant", "")})
        elif family == "SECTIONS":
            ok = _sections(c, data)
        elif family == "OUTLINE":
            ok = _outline(c, data)
        elif family in {"LAYERS", "HIERARCHY_NEST"}:
            ok = _layers(c, {"levels": [{"label": x} if isinstance(x, str) else x for x in data.get("levels", [])]})
        elif family == "SET_DIAGRAM":
            ok = _venn(c, data)
        elif family == "STAT_TILES":
            ok = _stat_tiles(c, data)
        elif family == "CAUSE_EFFECT":
            ok = _cause_effect(c, data)
        elif family == "FORMULA_CALLOUT":
            ok = _formula(c, data)
        elif family == "DEFINITION":
            ok = _definition(c, data)
        elif family == "QUOTE":
            ok = _quote(c, data)
        elif family == "FUNNEL":
            ok = _funnel(c, data)
        elif family in {"CHART_CATEGORICAL", "CHART_RANKING"}:
            ok = _chart(c, data, "bar")
        elif family == "CHART_TREND":
            ok = _chart(c, data, "line")
        elif family == "CHART_PART_TO_WHOLE":
            ok = _chart(c, data, "pie")
        elif family in {"FUNCTION_PLOT", "CALCULUS_VIZ"}:
            ok = _function_plot(c, data)
        elif family == "MIND_MAP":
            ok = _mind_map(c, data)
        elif family in {"KEY_POINTS", "LIST_STRUCTURED"}:
            cards = [(str(x.get("heading", "")), str(x.get("body", ""))) for x in (data.get("cards") or [])[:6] if isinstance(x, dict)]
            ok = len(cards) >= 2
            if ok:
                _grid(c, cards, icons=True)
        elif family == "CHECKLIST":
            items = [(str(x.get("text", x)) if isinstance(x, dict) else str(x), "") for x in (data.get("items") or [])[:8]]
            ok = len(items) >= 2
            if ok:
                _grid(c, [(t, "") for t, _ in items], numbered=True, cols=2 if len(items) > 4 else 1, icons=False)
        else:
            ok = False
        if not ok:  # any other family: a card grid from its text — still a visual, never bullets
            flat: List[str] = []
            _flatten(data, flat)
            flat = flat[:6]
            if len(flat) < 2:
                return False
            _grid(c, [(t[:60], t if len(t) > 60 else "") for t in flat], icons=True)
        if panel:
            _, size, label, paras = panel
            y0 = band.top + band.height - note_h
            with group(slide):
                c.shape(MSO_SHAPE.ROUNDED_RECTANGLE, band.left, y0, band.width, note_h, _blend(theme.primary_rgb, 0.93), None, radius=0.08)
                c.accent_bar(band.left, y0, note_h, 0.08)
                body = [P(label, size * 0.72, True, _readable_accent(theme))] if label else []
                body += [P(t, size, False, theme.text_rgb, 3) for t in paras]
                c.tb(band.left + 0.08 + PAD * 0.8, y0 + 0.08, band.width - 0.08 - PAD * 1.6, note_h - 0.16, body, MSO_ANCHOR.MIDDLE)
        return True
    except Exception as exc:  # best-effort: never break deck export
        logger.warning("pptx visual %s failed: %s", family, exc)
        return False
