"""
Slide chrome for the PPTX builder: header, footer and title slide, driven by the
deck's template (see ``templates.py``). Everything here is geometry on the same
0.5in margins as the rest of the deck; only the *treatment* changes per template.
"""

from __future__ import annotations

from typing import Callable, Optional

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from learnova.rendering import layout as L
from learnova.rendering.templates import get_template


# ── colour helpers (shared with ppt_builder) ──────────────────────────────────
def _lum(rgb) -> float:
    ch = []
    for v in (rgb[0], rgb[1], rgb[2]):
        c = v / 255.0
        ch.append(c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]


def _contrast(a, b) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _readable(bg, prefer=None) -> RGBColor:
    """Text colour that stays legible on ``bg``: ``prefer`` if it has contrast ≥ 4, else white/ink."""
    if prefer is not None and _contrast(prefer, bg) >= 4.0:
        return prefer
    return RGBColor(0xFF, 0xFF, 0xFF) if _contrast(RGBColor(255, 255, 255), bg) >= _contrast(RGBColor(0x11, 0x14, 0x22), bg) else RGBColor(0x11, 0x14, 0x22)


def _tint(rgb, t: float) -> RGBColor:
    return RGBColor(int(rgb[0] + (255 - rgb[0]) * t), int(rgb[1] + (255 - rgb[1]) * t), int(rgb[2] + (255 - rgb[2]) * t))


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=1.0, radius: Optional[float] = None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = min(0.5, radius)
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return tb


def _title_text(title: str, template) -> str:
    return title.upper() if template.title_case == "upper" else title


# ── header ────────────────────────────────────────────────────────────────────
def add_header(slide, title_text: str, theme) -> None:
    tpl = get_template(getattr(theme, "template_id", None))
    W, MX = L.SLIDE_W, L.MARGIN_X
    title = _title_text(title_text, tpl)
    full_w = W - 2 * MX

    if tpl.header == "bar":
        _rect(slide, 0, 0, W, 1.1, theme.primary_rgb)
        _rect(slide, 0, 1.1, W, 0.07, theme.accent_rgb)
        pt = L.fit_font_size([title], full_w, 0.85, max_pt=30, min_pt=15, bold=True)
        _text(slide, MX, 0.1, full_w, 0.9, title, pt, _readable(theme.primary_rgb, prefer=theme.accent_rgb), bold=True)
        return

    ink = _readable(theme.bg_rgb, prefer=theme.primary_rgb)
    if tpl.header == "underline":
        pt = L.fit_font_size([title], full_w, 0.85, max_pt=32, min_pt=15, bold=True)
        _text(slide, MX, 0.12, full_w, 0.86, title, pt, ink, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        _rect(slide, MX, 1.08, full_w, 0.02, _tint(theme.primary_rgb, 0.7))
        _rect(slide, MX, 1.05, 1.4, 0.07, theme.accent_rgb)
    elif tpl.header == "side":
        _rect(slide, 0, 0, 0.2, L.SLIDE_H, theme.primary_rgb)
        _rect(slide, 0.2, 0, 0.06, L.SLIDE_H, theme.accent_rgb)
        pt = L.fit_font_size([title], full_w - 0.2, 0.85, max_pt=30, min_pt=15, bold=True)
        _text(slide, MX + 0.1, 0.12, full_w - 0.1, 0.9, title, pt, ink, bold=True)
        _rect(slide, MX + 0.1, 1.1, full_w - 0.1, 0.015, _tint(theme.primary_rgb, 0.75))
    elif tpl.header == "ribbon":
        pt = L.fit_font_size([title], full_w - 0.8, 0.62, max_pt=26, min_pt=14, bold=True)
        est_w = min(full_w, max(3.2, len(title) * pt * 0.6 / 72 + 0.9))
        _rect(slide, MX, 0.25, est_w, 0.74, theme.primary_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, MX + 0.4, 0.25, est_w - 0.8, 0.74, title, pt, _readable(theme.primary_rgb, prefer=theme.accent_rgb), bold=True)
        _rect(slide, MX + est_w + 0.15, 0.55, 0.14, 0.14, theme.accent_rgb, MSO_SHAPE.OVAL)
    else:  # plain
        pt = L.fit_font_size([title], full_w, 0.85, max_pt=34, min_pt=16, bold=True)
        _text(slide, MX, 0.12, full_w, 0.86, title, pt, ink, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        _rect(slide, MX, 1.08, full_w, 0.012, _tint(theme.text_rgb, 0.8))


# ── footer ────────────────────────────────────────────────────────────────────
def add_footer(slide, n: int, total: int, theme, brand: str = "Learnova") -> None:
    tpl = get_template(getattr(theme, "template_id", None))
    W, MX = L.SLIDE_W, L.MARGIN_X
    w = 3.0
    if tpl.footer == "rule":
        _rect(slide, MX, L.FOOTER_Y - 0.04, W - 2 * MX, 0.01, _tint(theme.text_rgb, 0.82))
    left_x = MX + (0.1 if tpl.header == "side" else 0)
    _text(slide, left_x, L.FOOTER_Y, w, L.FOOTER_H, brand, 10, theme.subtext_rgb, bold=True)
    label = f"{n} / {total}"
    if tpl.footer == "pill":
        pw = 0.95
        _rect(slide, W - MX - pw, L.FOOTER_Y + 0.02, pw, L.FOOTER_H - 0.04, _tint(theme.primary_rgb, 0.86), MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, W - MX - pw, L.FOOTER_Y + 0.02, pw, L.FOOTER_H - 0.04, label, 10, _readable(_tint(theme.primary_rgb, 0.86), prefer=theme.primary_rgb), bold=True, align=PP_ALIGN.CENTER)
    else:
        _text(slide, W - MX - w, L.FOOTER_Y, w, L.FOOTER_H, label, 10, theme.subtext_rgb, align=PP_ALIGN.RIGHT)


# ── title slide ───────────────────────────────────────────────────────────────
def _decor(slide, theme, kind: str, on_dark: bool) -> None:
    """Background ornaments, drawn first so they sit behind the title."""
    W, H = L.SLIDE_W, L.SLIDE_H
    if kind == "circles":
        for d, col, x, y in ((6.2, _tint(theme.accent_rgb, 0.1 if on_dark else 0.8), 9.4, 3.9),
                             (3.4, _tint(theme.primary_rgb, 0.25 if on_dark else 0.85), 8.2, -1.4),
                             (1.6, _tint(theme.accent_rgb, 0.5 if on_dark else 0.7), 11.6, 0.6)):
            _rect(slide, x, y, d, d, col, MSO_SHAPE.OVAL)
    elif kind == "dots":
        col = _tint(theme.accent_rgb, 0.35 if on_dark else 0.4)
        for i in range(8):
            for j in range(5):
                _rect(slide, 0.6 + i * 0.28, 5.7 + j * 0.28, 0.07, 0.07, col, MSO_SHAPE.OVAL)
    elif kind == "lines":
        col = _tint(theme.accent_rgb if on_dark else theme.primary_rgb, 0.55 if on_dark else 0.75)
        for k in range(5):
            _rect(slide, 0.0, 6.2 + k * 0.22, W, 0.025, col)
    elif kind == "corner":
        _rect(slide, W - 3.0, H - 1.4, 3.0, 1.4, theme.accent_rgb)
        _rect(slide, W - 4.6, H - 0.7, 1.6, 0.7, _tint(theme.accent_rgb, 0.45))


def add_title_slide(prs, theme, topic_title: str, n_slides: int, hero_fn: Optional[Callable[[str], bytes]] = None, subtitle: str = ""):
    """``hero_fn(kind)`` returns PNG bytes for the vector illustration (``"panel"`` portrait or ``"banner"`` wide)."""
    tpl = get_template(getattr(theme, "template_id", None))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H, MX = L.SLIDE_W, L.SLIDE_H, L.MARGIN_X
    title = _title_text(topic_title, tpl)
    layout = tpl.title_layout
    sub_text = f"{subtitle}  •  {n_slides} slides" if subtitle else f"{n_slides} slides  •  made with Learnova"

    light = layout in {"centered", "left_band"}          # these layouts keep the light page colour
    page_bg = theme.bg_rgb if light else theme.primary_rgb
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = page_bg
    ink = _readable(page_bg, prefer=theme.primary_rgb if light else None)
    sub_ink = theme.subtext_rgb if light else _readable(page_bg, prefer=_tint(theme.primary_rgb, 0.7))

    def tsize(width, height, hi=54):
        base = hi if len(topic_title) < 28 else hi - 14 if len(topic_title) < 55 else hi - 22
        return L.fit_font_size([title], width, height, max_pt=base, min_pt=24, bold=True, spacing_pt=0)

    def picture(x, y, w, h, kind="panel"):
        if hero_fn is None:
            return
        try:
            from learnova.rendering.pptx_visuals import _Canvas
            _Canvas(slide, theme, L.content_band(has_takeaway=False)).picture(hero_fn(kind), x, y, w, h)
        except Exception:
            pass

    if layout == "hero_split":
        _decor(slide, theme, tpl.decoration, on_dark=True)
        picture(8.85, 0.95, 3.95, 4.94)
        _rect(slide, 0.9, 2.0, 1.4, 0.14, theme.accent_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, 0.8, 2.3, 7.8, 2.6, title, tsize(7.8, 2.6), ink, bold=True, anchor=MSO_ANCHOR.TOP)
        _text(slide, 0.9, 5.0, 7.6, 0.5, sub_text, 18, sub_ink)
    elif layout == "centered":
        _rect(slide, 0, 0, W, 0.35, theme.primary_rgb)
        _decor(slide, theme, tpl.decoration, on_dark=False)
        _rect(slide, W / 2 - 0.7, 1.15, 1.4, 0.12, theme.accent_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, 1.2, 1.4, W - 2.4, 1.9, title, tsize(W - 2.4, 1.9, 50), ink, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, 1.2, 3.35, W - 2.4, 0.45, sub_text, 18, sub_ink, align=PP_ALIGN.CENTER)
        picture(W / 2 - 2.4, 4.0, 4.8, 2.98, "banner")
    elif layout == "left_band":
        band_w = 4.4
        _rect(slide, 0, 0, band_w, H, theme.primary_rgb)
        _rect(slide, band_w, 0, 0.1, H, theme.accent_rgb)
        picture(0.45, 1.4, 3.5, 4.4)
        _text(slide, band_w + 0.9, 2.0, W - band_w - 1.5, 2.6, title, tsize(W - band_w - 1.5, 2.6, 46), ink, bold=True, anchor=MSO_ANCHOR.TOP)
        _rect(slide, band_w + 0.9, 4.9, 1.2, 0.1, theme.accent_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, band_w + 0.9, 5.15, W - band_w - 1.5, 0.5, sub_text, 18, sub_ink)
        if tpl.decoration == "dots":
            _decor(slide, theme, "dots", on_dark=False)
    elif layout == "frame":
        _rect(slide, 0.35, 0.35, W - 0.7, H - 0.7, None, line=theme.accent_rgb, line_w=2.0)
        _rect(slide, 0.5, 0.5, W - 1.0, H - 1.0, None, line=_tint(theme.primary_rgb, 0.5), line_w=0.75)
        _decor(slide, theme, tpl.decoration, on_dark=True)
        picture(8.7, 1.1, 3.75, 4.7)
        _rect(slide, 1.0, 2.05, 1.2, 0.1, theme.accent_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, 1.0, 2.35, 7.2, 2.5, title, tsize(7.2, 2.5, 48), ink, bold=True, anchor=MSO_ANCHOR.TOP)
        _text(slide, 1.0, 5.0, 7.2, 0.5, sub_text, 18, sub_ink)
    else:  # poster
        _decor(slide, theme, tpl.decoration, on_dark=True)
        _rect(slide, 0, 0, 0.35, H, theme.accent_rgb)
        _text(slide, 0.9, 1.5, W - 4.6, 3.8, title, tsize(W - 4.6, 3.8, 66), ink, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _rect(slide, 0.9, 5.5, 2.2, 0.14, theme.accent_rgb, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        _text(slide, 0.9, 5.8, 8.0, 0.5, sub_text, 18, sub_ink)
        picture(W - 3.9, 0.6, 3.1, 3.9)
    return slide
