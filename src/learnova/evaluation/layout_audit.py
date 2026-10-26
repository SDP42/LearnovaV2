"""
Layout audit for generated PPTX decks: whitespace, sparse cards, overflow, fonts, grid.

Works on the shapes themselves (no rendering), so it takes milliseconds and can run
in tests and in the pipeline. Every slide gets measurable numbers instead of an
eyeball judgement:

* ``coverage``      share of the content band covered by something visible (cards, text, pictures)
* ``bottom_gap``    empty inches between the last element and the bottom of the band
* ``sparse_cards``  filled boxes whose text occupies <30% of their height (text floating in a big box)
* ``overflow``      elements past the slide edge or the 0.5in margins
* ``fonts``/``min_pt``  typefaces and the smallest size used
* ``edges``         distinct text left-edges (a tidy grid has few)

    python -m learnova.evaluation.layout_audit deck.pptx
"""

from __future__ import annotations

import io
import sys
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Emu

from learnova.rendering import layout as L

_BAND_TOP = L.CONTENT_TOP
_BAND_BOTTOM = L.SLIDE_H - L.BOTTOM_MARGIN
_GRID_X, _GRID_Y = 80, 48


def _in(v) -> float:
    return Emu(int(v or 0)).inches


def _runs_pt(shape) -> List[float]:
    sizes: List[float] = []
    if not shape.has_text_frame:
        return sizes
    for p in shape.text_frame.paragraphs:
        if p.font and p.font.size:
            sizes.append(p.font.size.pt)
        for r in p.runs:
            if r.font and r.font.size:
                sizes.append(r.font.size.pt)
    return sizes


def _fonts(shape) -> List[str]:
    names: List[str] = []
    if not shape.has_text_frame:
        return names
    for p in shape.text_frame.paragraphs:
        if p.font and p.font.name:
            names.append(p.font.name)
        for r in p.runs:
            if r.font and r.font.name:
                names.append(r.font.name)
    return names


def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:                      # group
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _visible(shape) -> bool:
    if shape.shape_type == 13 or getattr(shape, "has_table", False) or getattr(shape, "has_chart", False):   # picture, table, chart
        return True
    try:
        if shape.fill.type == 1:                    # solid
            return True
    except Exception:
        pass
    try:
        if shape.line.fill.type == 1:
            return True
    except Exception:
        pass
    return bool(shape.has_text_frame and shape.text_frame.text.strip())


def _text_fill(shape) -> float:
    """Estimated share of a box's height its text occupies (1.0 = full)."""
    tf = shape.text_frame
    texts = [p.text for p in tf.paragraphs if p.text.strip()]
    if not texts:
        return 1.0
    sizes = _runs_pt(shape) or [14.0]
    pt = max(sizes)
    w = max(0.3, _in(shape.width) - 0.3)
    h = max(0.1, _in(shape.height))
    return min(1.0, L.block_height(texts, w, pt, spacing_pt=4.0) / h)


def _broken_words(shape) -> List[str]:
    """Words wider than their text box at their own size — they would be cut mid-word ("Quadrati/c")."""
    if not shape.has_text_frame or not shape.text_frame.word_wrap:
        return []
    tf = shape.text_frame
    room = _in(shape.width) - _in(tf.margin_left or 0) - _in(tf.margin_right or 0)
    out: List[str] = []
    for p in tf.paragraphs:
        for r in p.runs:
            pt = r.font.size.pt if r.font and r.font.size else 18.0
            for piece in L._pieces(r.text):
                if L.word_width_in(piece, pt, bool(r.font and r.font.bold)) > room + 0.01:
                    out.append(piece)
    return out


def audit_slide(slide, index: int) -> Dict[str, Any]:
    grid = [[False] * _GRID_X for _ in range(_GRID_Y)]
    bottoms: List[float] = []
    sparse: List[Dict[str, Any]] = []
    overflow: List[str] = []
    broken: List[str] = []
    sizes: List[float] = []
    fonts = set()
    edges = set()
    words = 0

    for sh in _iter_shapes(slide.shapes):
        left, top = _in(sh.left), _in(sh.top)
        w, h = _in(sh.width), _in(sh.height)
        right, bottom = left + w, top + h
        sizes += _runs_pt(sh)
        fonts.update(_fonts(sh))
        broken += _broken_words(sh)
        if sh.has_text_frame and sh.text_frame.text.strip():
            words += len(sh.text_frame.text.split())
            if top >= _BAND_TOP - 0.05:
                edges.add(round(left * 20) / 20)
        # header / footer chrome is outside the content band
        if bottom <= _BAND_TOP + 0.05 or top >= _BAND_BOTTOM + 0.05:
            continue
        if w >= L.SLIDE_W - 0.2 and h >= L.SLIDE_H - 0.2:
            continue                                # full-slide background
        if not _visible(sh) or sh.name.startswith("Oval"):         # decorative circles may bleed off the slide
            continue
        if right > L.SLIDE_W + 0.02 or bottom > L.SLIDE_H + 0.02 or left < -0.02:
            overflow.append(f"{sh.name}: off the slide")
        elif right > L.SLIDE_W - L.MARGIN_X + 0.15 and w < L.SLIDE_W - 1:
            overflow.append(f"{sh.name}: past the right margin")
        x0, x1 = max(0.0, left), min(L.SLIDE_W, right)
        y0, y1 = max(_BAND_TOP, top), min(_BAND_BOTTOM, bottom)
        if x1 > x0 and y1 > y0:
            for gy in range(int((y0 - _BAND_TOP) / (_BAND_BOTTOM - _BAND_TOP) * _GRID_Y),
                            min(_GRID_Y, int((y1 - _BAND_TOP) / (_BAND_BOTTOM - _BAND_TOP) * _GRID_Y) + 1)):
                for gx in range(int(x0 / L.SLIDE_W * _GRID_X), min(_GRID_X, int(x1 / L.SLIDE_W * _GRID_X) + 1)):
                    grid[gy][gx] = True
            bottoms.append(y1)
        if sh.has_text_frame and sh.text_frame.text.strip() and h > 1.0 and w > 1.2:
            try:
                filled = sh.fill.type == 1
            except Exception:
                filled = False
            if filled:
                f = _text_fill(sh)
                if f < 0.3:
                    sparse.append({"shape": sh.name, "height_in": round(h, 2), "text_fill": round(f, 2)})

    if words <= 6 and any("thank you" in (sh.text_frame.text.lower() if sh.has_text_frame else "") for sh in slide.shapes):
        return {"index": index, "coverage": 1.0, "bottom_gap": 0.0, "words": words, "sparse_cards": [], "overflow": [],
                "min_pt": min(sizes) if sizes else None, "max_pt": max(sizes) if sizes else None, "fonts": sorted(fonts),
                "edges": sorted(edges), "flags": [], "kind": "closing"}
    covered = sum(sum(row) for row in grid)
    coverage = covered / float(_GRID_X * _GRID_Y)
    gap = round(_BAND_BOTTOM - max(bottoms), 2) if bottoms else round(_BAND_BOTTOM - _BAND_TOP, 2)
    flags: List[str] = []
    if coverage < 0.40 and words < 120:
        flags.append("sparse")
    if gap > 1.4:
        flags.append("empty bottom")
    if sparse:
        flags.append("floating text in big box")
    if overflow:
        flags.append("overflow")
    if broken:
        flags.append("word breaks: " + ", ".join(sorted(set(broken))[:4]))
    return {
        "index": index, "coverage": round(coverage, 2), "bottom_gap": gap, "words": words,
        "sparse_cards": sparse, "overflow": overflow, "min_pt": min(sizes) if sizes else None,
        "max_pt": max(sizes) if sizes else None, "fonts": sorted(fonts), "edges": sorted(edges), "flags": flags,
    }


def audit_pptx(data: bytes) -> Dict[str, Any]:
    prs = Presentation(io.BytesIO(data))
    slides = [audit_slide(s, i) for i, s in enumerate(prs.slides)]
    body = slides[1:] or slides                      # slide 0 is the title slide
    fonts = sorted({f for s in body for f in s["fonts"]})
    bad = [s for s in body if s["flags"]]
    return {
        "slides": len(slides),
        "mean_coverage": round(sum(s["coverage"] for s in body) / max(1, len(body)), 2),
        "flagged": len(bad),
        "flagged_pct": round(100 * len(bad) / max(1, len(body))),
        "fonts": fonts,
        "min_pt": min((s["min_pt"] for s in body if s["min_pt"]), default=None),
        "per_slide": slides,
    }


def main(argv: List[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    with open(argv[0], "rb") as fh:
        rep = audit_pptx(fh.read())
    print(f"{rep['slides']} slides · mean coverage {rep['mean_coverage']:.0%} · flagged {rep['flagged']} ({rep['flagged_pct']}%) "
          f"· fonts {rep['fonts']} · smallest text {rep['min_pt']}pt")
    for s in rep["per_slide"][1:]:
        if s["flags"]:
            print(f"  slide {s['index'] + 1}: {', '.join(s['flags'])} · coverage {s['coverage']:.0%} · bottom gap {s['bottom_gap']}in · {s['words']} words")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
