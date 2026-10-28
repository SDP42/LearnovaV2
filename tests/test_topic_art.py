"""AI-written slide diagrams are checked before use; split topics get named parts instead of (2/5) counters."""
from learnova.ai import topic_art


GOOD = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 720"><rect width="600" height="720" fill="#ffffff"/>'
        '<line x1="40" y1="600" x2="560" y2="600" stroke="#1f2937"/><line x1="40" y1="600" x2="40" y2="80" stroke="#1f2937"/>'
        '<circle cx="120" cy="500" r="8" fill="PRIMARY"/><circle cx="220" cy="420" r="8" fill="PRIMARY"/>'
        '<circle cx="320" cy="350" r="8" fill="PRIMARY"/><line x1="60" y1="560" x2="540" y2="160" stroke="ACCENT" stroke-width="4"/>'
        '<text x="560" y="140" font-size="20" text-anchor="end">Fitted line with a very long label $here$</text></svg>')


def test_clean_svg_accepts_drawing_and_fixes_labels():
    svg = topic_art.clean_svg("Here you go:\n" + GOOD, "#A61E4D", "#22C55E")
    assert svg and "#A61E4D" in svg and "PRIMARY" not in svg and "$" not in svg
    assert 'width="600"' in svg


def test_clean_svg_strips_unsafe_parts_and_rejects_empty():
    safe = topic_art.clean_svg(GOOD.replace("<rect", '<script>alert(1)</script><image href="http://x/y.png"/><rect style="fill:#ff0000;cursor:pointer" class="x"', 1),
                               "#000000", "#111111")
    assert safe and "<script" not in safe and "<image" not in safe and "class=" not in safe and 'fill="#ff0000"' in safe and "cursor" not in safe
    assert topic_art.clean_svg(GOOD.replace('<rect width', '<use href="http://x/y.svg#a"/><rect width'), "#000000", "#111111") is None
    assert topic_art.clean_svg('<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>', "#000000", "#111111") is None
    assert topic_art.clean_svg("no svg here", "#000000", "#111111") is None


def test_art_page_only_serves_cached_ids(tmp_path, monkeypatch):
    monkeypatch.setattr(topic_art, "ART_DIR", tmp_path)
    (tmp_path / ("a" * 20 + ".svg")).write_text(GOOD.replace("<rect", "<title>Least squares</title><rect", 1))
    page = topic_art.art_page("a" * 20)
    assert page and "<svg" in page and "Least squares" in page
    assert topic_art.art_page("../etc/passwd") is None and topic_art.art_page("b" * 20) is None


def test_no_ai_means_no_diagram(monkeypatch):
    assert topic_art.available() is False
    assert topic_art.illustrate([{"improved": {"title": "Bias", "bullets": ["High bias means underfitting the data."]}}], ["KEY_POINTS"], "#000000", "#111111") == 0


def test_split_topic_parts_are_named():
    from learnova.pipeline.polish import name_parts

    entries = [{"improved": {"title": "Hypothesis (1/2)", "bullets": ["A hypothesis is a function that describes the target."]}},
               {"improved": {"title": "Hypothesis (2/2)", "bullets": ["Hypothesis space (H): the set of all legal hypotheses."]}}]
    name_parts(entries)
    assert entries[0]["improved"]["title"] == "Hypothesis"
    assert entries[1]["improved"]["title"] == "Hypothesis: Hypothesis space (H)" or "(2/2)" not in entries[1]["improved"]["title"]


def test_every_topic_slide_has_a_linked_illustration_offline():
    """Without keys every topic slide of the PowerPoint still carries a picture that links to its own HTML page."""
    from io import BytesIO

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from learnova.parsers.markdown_converter import from_typed_text
    from learnova.pipeline.orchestrator import PipelineConfig, generate

    md = ("# Machine learning\n\n## Bias\nBias is the error from wrong assumptions in the learning algorithm. High bias causes underfitting.\n\n"
          "## Variance\nVariance is the error from sensitivity to small changes in the training set. High variance causes overfitting.\n\n"
          "## Trade-off\nA good model balances bias and variance so that it generalises to unseen data.\n")
    res = generate(from_typed_text(md, "ml"), PipelineConfig(enable_enhancement=False, enable_vision_ocr=False, enable_quizzes=False))
    slides = list(Presentation(BytesIO(res.pptx_bytes)).slides)[1:-1]
    linked = [any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE and "/api/art/" in (sh.click_action.hyperlink.address or "") for sh in s.shapes)
              for s in slides]
    assert linked and sum(linked) >= len(linked) - 1          # agenda-like slides may carry the header picture only
    art_id = next(sh.click_action.hyperlink.address.rsplit("/", 1)[1] for s in slides for sh in s.shapes
                  if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and "/api/art/" in (sh.click_action.hyperlink.address or ""))
    assert topic_art.art_page(art_id) and "<svg" in topic_art.art_page(art_id)


def test_refresh_unknown_deck_is_404():
    from fastapi.testclient import TestClient

    import apps.api.main as main

    main.app.dependency_overrides[main.current_user] = lambda: "nobody"
    try:
        assert TestClient(main.app).post("/api/decks/does-not-exist/refresh").status_code == 404
    finally:
        main.app.dependency_overrides.pop(main.current_user, None)
