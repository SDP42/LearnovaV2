"""Interactive labs: found by topic, embedded in the web deck, linked from the PowerPoint, served and uploadable via the API."""
import pytest
from fastapi.testclient import TestClient

from learnova import labs
from learnova.parsers.markdown_converter import from_typed_text
from learnova.pipeline.orchestrator import PipelineConfig, generate

_MD = ("# Segmentation\n\n## Region growing\nRegion growing groups pixels into larger regions based on a condition. "
       "We start from a seed pixel and add neighbours while max minus min stays within a threshold T.\n\n"
       "## Thresholding\nA single threshold separates the object from the background by intensity.\n")


def test_lab_is_found_by_topic_only():
    assert labs.find_lab("Region growing")["id"] == "region_growing"
    assert labs.find_lab("Region based segmentation (2/3)")["id"] == "region_growing"
    assert labs.find_lab("Edge detection") is None


def test_deck_embeds_and_links_the_lab():
    res = generate(from_typed_text(_MD, "seg"), PipelineConfig(enable_enhancement=False, enable_vision_ocr=False, enable_quizzes=False))
    html = res.html_bytes.decode("utf-8")
    assert html.count('class="lv-labbtn"') == 1 and 'id="lab-region_growing-doc"' in html
    from io import BytesIO

    from pptx import Presentation

    links = [sh.click_action.hyperlink.address for s in Presentation(BytesIO(res.pptx_bytes)).slides for sh in s.shapes
             if sh.has_text_frame and "interactive lab" in sh.text_frame.text]
    assert links == [labs.lab_url("region_growing")]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    import apps.api.main as main

    monkeypatch.setattr(labs, "USER_DIR", tmp_path)
    main.app.dependency_overrides[main.current_user] = lambda: "u1"
    yield TestClient(main.app)
    main.app.dependency_overrides.pop(main.current_user, None)


def test_api_serves_and_accepts_labs(client):
    assert client.get("/api/labs/region_growing").status_code == 200
    assert client.get("/api/labs/../secret").status_code == 404
    r = client.post("/api/labs", json={"id": "otsu", "topics": ["otsu thresholding"], "html": "<html><head><title>Otsu Lab</title></head><body>x</body></html>"})
    assert r.status_code == 201 and r.json()["name"] == "Otsu Lab"
    assert any(l["id"] == "otsu" for l in client.get("/api/labs").json()["labs"])
    assert labs.find_lab("Otsu thresholding")["id"] == "otsu"
    assert client.post("/api/labs", json={"id": "Bad Id!", "topics": ["x"], "html": "<html></html>"}).status_code == 400
