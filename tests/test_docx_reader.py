import docx

from learnova.parsers.markdown_converter import docx_to_markdown


def test_word_file_becomes_sections_lists_and_tables(tmp_path):
    d = docx.Document()
    d.add_heading("Word sense disambiguation", level=2)
    d.add_paragraph("A word can have several senses.")
    d.add_paragraph("Supervised methods", style="List Bullet")
    d.add_paragraph("Dictionary methods", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Method", "Needs", "Lesk", "a dictionary"
    path = tmp_path / "w.docx"
    d.save(path)
    md = docx_to_markdown(str(path))
    assert "## Word sense disambiguation" in md            # top heading level → section
    assert "- Supervised methods" in md and "| Lesk | a dictionary |" in md
