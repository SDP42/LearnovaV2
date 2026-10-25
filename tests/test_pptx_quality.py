"""PPTX export quality: schema order of transition/timing, reopen-ability, animation groups, bounds."""

from __future__ import annotations

import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from learnova.rendering.ppt_builder import build_pptx
from learnova.rendering.pptx_animation import apply_click_builds, set_transition

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _deck():
    cards = [{"heading": f"Idea {i}", "body": f"A fairly descriptive sentence about idea number {i}."} for i in range(1, 5)]
    out = []
    for i in range(1, 5):
        out.append({"original": {"text": "x"}, "improved": {
            "layout_type": "MINIMAL_TEXT", "title": f"Topic {i}",
            "bullets": ["First point about the topic", "Second point about the topic", "Third point"],
            "takeaway": "Remember this."}})
    out.append({"original": {"text": "x"}, "improved": {
        "layout_type": "QUIZ", "title": "Checkpoint", "question": "Which one?",
        "options": ["A) one", "B) two", "C) three", "D) four"], "correct": "B", "explanation": "Because two."}})
    return out


@pytest.fixture(scope="module")
def prs():
    return Presentation(io.BytesIO(build_pptx(_deck(), "Quality Deck", theme_id="auto")))


def _order(slide):
    return [etree.QName(c).localname for c in slide._element]


def test_reopens_and_has_all_slides(prs):
    assert len(prs.slides) == 7  # title + 5 + closing


def test_transition_before_timing_after_clrmapovr(prs):
    seen = 0
    for slide in prs.slides:
        order = _order(slide)
        if "transition" in order:
            seen += 1
            assert order.index("clrMapOvr") < order.index("transition")
            if "timing" in order:
                assert order.index("transition") < order.index("timing")
        if "timing" in order:
            assert order.index("clrMapOvr") < order.index("timing")
    assert seen >= 1


def test_title_slide_has_no_transition(prs):
    assert "transition" not in _order(prs.slides[0])


@pytest.mark.parametrize("name,tag", [("fade", "fade"), ("slide", "push"), ("convex", "cover"),
                                      ("concave", "pull"), ("zoom", "zoom")])
def test_transition_mapping(name, tag):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    assert set_transition(s, name)
    tr = s._element.find(f"{{{P}}}transition")
    assert tr.get("spd") == "med"
    assert etree.QName(tr[0]).localname == tag
    assert _order(s).index("transition") > _order(s).index("clrMapOvr")


def test_none_transition_writes_nothing():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    assert not set_transition(s, "none")
    assert "transition" not in _order(s)


def test_timing_goes_after_transition_when_both_set():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    a = s.shapes.add_textbox(0, 0, 100, 100)
    b = s.shapes.add_textbox(0, 0, 100, 100)
    set_transition(s, "fade")
    assert apply_click_builds(s, [[a.shape_id], [b.shape_id]])
    set_transition(s, "slide")  # re-setting must keep the order
    o = _order(s)
    assert o.index("clrMapOvr") < o.index("transition") < o.index("timing")


def test_groups_recorded_and_ids_exist():
    out = build_pptx(_deck(), "Quality Deck", theme_id="auto")
    p = Presentation(io.BytesIO(out))
    quiz = p.slides[5]
    ids = {sh.shape_id for sh in quiz.shapes}
    timing = quiz._element.find(f"{{{P}}}timing")
    assert timing is not None
    targets = {int(t.get("spid")) for t in timing.iter(f"{{{P}}}spTgt")}
    assert targets and targets <= ids


def test_animation_opt_out(monkeypatch):
    monkeypatch.setenv("LEARNOVA_PPTX_ANIM", "0")
    p = Presentation(io.BytesIO(build_pptx(_deck(), "Quality Deck", theme_id="auto")))
    assert all("timing" not in _order(s) for s in p.slides)


def test_no_shape_leaves_the_slide(prs):
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    for i, slide in enumerate(prs.slides):
        if i in (0, len(prs.slides) - 1):   # decorative circles intentionally bleed off the title/closing slides
            continue
        for sh in slide.shapes:
            r, b = Emu(sh.left + sh.width).inches, Emu(sh.top + sh.height).inches
            if sh.top is not None and Emu(sh.top).inches < -0.01 or Emu(sh.left).inches < -0.01:
                continue  # header decoration
            assert r <= W + 0.02 and b <= H + 0.02, (i, sh.name, r, b)
