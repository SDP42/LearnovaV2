"""Copies each slide's part of the speaker script into that slide's PowerPoint notes.

    python scripts/research/add_speaker_notes.py docs/review/speaker_script.md docs/review/Learnova_SIH_Review.pptx
"""
import re
import sys

from pptx import Presentation

md, deck = sys.argv[1], sys.argv[2]
parts = re.split(r"^# ", open(md, encoding="utf-8").read(), flags=re.M)
notes = {}
for part in parts:
    m = re.match(r"Slide (\d+) — (.*)\n", part)
    if m:
        body = part[m.end():].strip()
        body = re.sub(r"\*\*(.+?)\*\*", r"\1", body)
        notes[int(m.group(1))] = f"{m.group(2)}\n\n{body}"
prs = Presentation(deck)
for i, slide in enumerate(prs.slides, 1):
    if i in notes:
        slide.notes_slide.notes_text_frame.text = notes[i]
prs.save(deck)
print("notes on", len(notes), "slides")
