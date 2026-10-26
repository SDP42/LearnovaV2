"""
End-to-end health check of the API, in-process (no server needed).

    python scripts/e2e_check.py            # all scenarios
    python scripts/e2e_check.py typed pdf  # some scenarios

Signs in through the dev OTP, then drives the real routes: typed text, a short prompt, markdown/PDF/PPTX
uploads, generation, deck/analysis/download/thumb/tutor, plus the error paths. Prints one line per check and a
summary; exits non-zero when anything fails.
"""
import io
import os
import pathlib
import sys
import time

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))
os.environ.setdefault("LEARNOVA_RATE_LIMIT", "off")
os.environ.setdefault("OTP_DEV_MODE", "1")
os.environ.setdefault("LEARNOVA_DB_PATH", str(ROOT / ".data" / "e2e.db"))
import logging  # noqa: E402

logging.disable(logging.CRITICAL)
from fastapi.testclient import TestClient  # noqa: E402

from apps.api.main import app  # noqa: E402

client = TestClient(app)
results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(("  ok   " if ok else "  FAIL ") + name + (f"  — {detail}" if detail else ""))


def login():
    phone = f"+9190{int(time.time()) % 100000000:08d}"
    r = client.post("/api/auth/otp/request", json={"phone": phone}).json()
    v = client.post("/api/auth/otp/verify", json={"phone": phone, "code": r["dev_otp"]}).json()
    return {"Authorization": f"Bearer {v['token']}"}


def wait(job_id, h, timeout=240, want=("awaiting_review", "done", "failed")):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = client.get(f"/api/jobs/{job_id}", headers=h).json()
        if j["status"] in want:
            return j, time.time() - t0
        time.sleep(0.4)
    return client.get(f"/api/jobs/{job_id}", headers=h).json(), time.time() - t0


def generate(job_id, h, **opts):
    body = {"enable_quizzes": True, "enable_enhancement": False, "enable_vision_ocr": False, **opts}
    r = client.post(f"/api/jobs/{job_id}/generate", json=body, headers=h)
    if r.status_code != 202:
        return None, 0, r
    j, dt = wait(job_id, h, want=("done", "failed"))
    return j, dt, r


def verify_deck(label, job_id, h, min_slides=1):
    d = client.get(f"/api/jobs/{job_id}/deck", headers=h)
    slides = d.json().get("slides", []) if d.status_code == 200 else []
    check(f"{label}: deck has slides", len(slides) >= min_slides, f"{len(slides)} slides")
    for art in ("pptx", "html"):
        r = client.get(f"/api/jobs/{job_id}/download/{art}", headers=h)
        check(f"{label}: download {art}", r.status_code == 200 and len(r.content) > 5000, f"{r.status_code} {len(r.content)}B")
    a = client.get(f"/api/decks/{job_id}/analysis", headers=h)
    check(f"{label}: analysis", a.status_code == 200 and "summary" in a.json())
    t = client.get(f"/api/decks/{job_id}/thumb", headers=h)
    check(f"{label}: thumbnail", t.status_code == 200 and len(t.content) > 500, f"{t.status_code}")
    c = client.post(f"/api/decks/{job_id}/chat", json={"message": "summarise", "persona": "teacher"}, headers=h)
    check(f"{label}: tutor", c.status_code == 200 and c.json().get("answer"))
    e = client.get(f"/api/decks/{job_id}/editable", headers=h)
    check(f"{label}: editable slides", e.status_code == 200)
    return slides


SAMPLE = (ROOT / "demo" / "sample_syllabus.md").read_text()
SHORT = "Colour theory for designers"          # covered by the built-in lesson library, so it works without an AI key
UNKNOWN = "Quantum error correction codes"     # not in the library: needs an AI key


def sc_typed(h):
    print("typed syllabus")
    r = client.post("/api/jobs/typed", json={"text": SAMPLE, "title": "Typed Syllabus"}, headers=h)
    check("typed: created", r.status_code == 201, str(r.status_code))
    jid = r.json()["id"]
    j, dt, _ = generate(jid, h)
    check("typed: generated", j and j["status"] == "done", f"{dt:.1f}s {j and j.get('error')}")
    if j and j["status"] == "done":
        verify_deck("typed", jid, h, min_slides=5)


def sc_short(h):
    print("short prompt")
    r = client.post("/api/jobs/typed", json={"text": SHORT, "title": SHORT}, headers=h)
    jid = r.json()["id"]
    j, dt, _ = generate(jid, h)
    check("short: generated", j and j["status"] == "done", f"{dt:.1f}s {j and j.get('error')}")
    if j and j["status"] == "done":
        verify_deck("short", jid, h, min_slides=5)


def sc_unknown(h):
    print("short prompt, nothing to match")
    r = client.post("/api/jobs/typed", json={"text": UNKNOWN, "title": UNKNOWN}, headers=h).json()
    ex = r.get("expansion") or {}
    llm = ex.get("mode") == "llm"
    check("unknown prompt: told what to do", ex.get("mode") in {"llm", "none"} and (llm or "AI key" in ex.get("message", "")), str(ex)[:90])


def sc_md(h):
    print("markdown upload")
    r = client.post("/api/jobs", files={"file": ("notes.md", (ROOT / "demo" / "sample_ml.md").read_bytes(), "text/markdown")}, headers=h)
    check("md: accepted", r.status_code == 202, str(r.status_code))
    jid = r.json()["id"]
    j, dt = wait(jid, h)
    check("md: converted", j["status"] == "awaiting_review", f"{dt:.1f}s {j.get('error')}")
    j, dt, _ = generate(jid, h)
    check("md: generated", j and j["status"] == "done", f"{dt:.1f}s")


def sc_pdf(h):
    print("pdf upload")
    pdf = pathlib.Path(os.environ.get("E2E_PDF", ""))
    if not pdf.is_file():
        print("  skip (set E2E_PDF=/path/to.pdf)")
        return
    r = client.post("/api/jobs", files={"file": (pdf.name, pdf.read_bytes(), "application/pdf")}, headers=h)
    check("pdf: accepted", r.status_code == 202)
    jid = r.json()["id"]
    j, dt = wait(jid, h)
    check("pdf: converted", j["status"] == "awaiting_review", f"{dt:.1f}s {j.get('error')}")
    j, dt, _ = generate(jid, h)
    check("pdf: generated", j and j["status"] == "done", f"{dt:.1f}s")
    if j and j["status"] == "done":
        slides = verify_deck("pdf", jid, h)
        try:
            import fitz

            pages = len(fitz.open(os.environ["E2E_PDF"]))
        except Exception:
            pages = 29
        check("pdf: slide count sane", len(slides) <= pages * 2.5, f"{len(slides)} slides")


def sc_pptx(h):
    print("pptx upload")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, (t, b) in enumerate([("Photosynthesis", "Plants convert light, water and CO2 into glucose and oxygen."),
                                ("Stages", "Light reactions happen in the thylakoid. The Calvin cycle happens in the stroma."),
                                ("Why it matters", "It feeds food chains and produces oxygen.")]):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = b
    buf = io.BytesIO()
    prs.save(buf)
    r = client.post("/api/jobs", files={"file": ("bio.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, headers=h)
    check("pptx: accepted", r.status_code == 202)
    jid = r.json()["id"]
    j, dt = wait(jid, h)
    check("pptx: converted", j["status"] == "awaiting_review", f"{dt:.1f}s {j.get('error')}")
    j, dt, _ = generate(jid, h)
    check("pptx: generated", j and j["status"] == "done", f"{dt:.1f}s")


def sc_errors(h):
    print("error paths")
    check("empty text rejected", client.post("/api/jobs/typed", json={"text": "  ", "title": "x"}, headers=h).status_code == 400)
    check("bad suffix rejected", client.post("/api/jobs", files={"file": ("x.exe", b"MZ", "application/octet-stream")}, headers=h).status_code == 400)
    r = client.post("/api/jobs", files={"file": ("broken.pdf", b"%PDF-1.4 not really", "application/pdf")}, headers=h)
    jid = r.json()["id"]
    j, _ = wait(jid, h)
    check("corrupt pdf fails cleanly", j["status"] == "failed" and j.get("error"), f"{j['status']}: {str(j.get('error'))[:60]}")
    check("unknown job 404", client.get("/api/jobs/nope", headers=h).status_code == 404)
    check("no token 401", client.get("/api/jobs").status_code == 401)
    r = client.post("/api/jobs/typed", json={"text": "x" * 400_000, "title": "huge"}, headers=h)
    check("huge text handled", r.status_code in (201, 400, 413), str(r.status_code))
    other = login()
    jid = client.post("/api/jobs/typed", json={"text": "# A\nSome text here that is long enough to chunk.", "title": "mine"}, headers=h).json()["id"]
    check("other user cannot read job", client.get(f"/api/jobs/{jid}", headers=other).status_code in (403, 404))


def sc_concurrent(h):
    print("three jobs at once")
    ids = []
    for i in range(3):
        jid = client.post("/api/jobs/typed", json={"text": SAMPLE, "title": f"concurrent {i}"}, headers=h).json()["id"]
        client.post(f"/api/jobs/{jid}/generate", json={"enable_enhancement": False, "enable_vision_ocr": False}, headers=h)
        ids.append(jid)
    t0 = time.time()
    done = 0
    for jid in ids:
        j, _ = wait(jid, h, want=("done", "failed"))
        done += j["status"] == "done"
    check("concurrent: all finished", done == 3, f"{done}/3 in {time.time() - t0:.1f}s")


SCENARIOS = {"typed": sc_typed, "short": sc_short, "unknown": sc_unknown, "md": sc_md, "pdf": sc_pdf, "pptx": sc_pptx, "errors": sc_errors, "concurrent": sc_concurrent}

if __name__ == "__main__":
    which = sys.argv[1:] or list(SCENARIOS)
    h = login()
    for name in which:
        t0 = time.time()
        try:
            SCENARIOS[name](h)
        except Exception as exc:  # a crash is a failure, with its message
            check(f"{name}: crashed", False, f"{type(exc).__name__}: {exc}")
        print(f"  ({time.time() - t0:.1f}s)")
    bad = [r for r in results if not r[1]]
    print(f"\n{len(results) - len(bad)}/{len(results)} checks passed")
    sys.exit(1 if bad else 0)
